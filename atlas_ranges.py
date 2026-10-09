"""Layers: anything with a footprint on a map (creatures, plants, characters,
paths, quests), painted over a map sheet.

Each workbook says how its world keeps time and whether it rates how common
things are (Format -> Layer times and rarity):

- time: none (things are just there), Day / Night, Dawn / Day / Dusk / Night,
  hourly, the Biomes six ticks, or a list of your own periods, optionally
  grouped (e.g. Day = 08, 12, 16)
- rarity: off (a square either has the thing or not), or up to three named
  levels, most common first

A layer records, per map and per period, which squares it's in (and at what
level, when rarity is on). On disk it's a plain sheet called "Ranges"
(Layer | Map | Times | Abundance | Squares | Colour | Style), with squares
written the way the Gazetteer writes them ("I6", "G9:K12"). Atlas rewrites
that sheet on save.
"""

import colorsys
import re

from PySide6.QtCore import QPoint, QPointF, QRect, QRectF, QSize, Qt, Signal
from PySide6.QtGui import QBrush, QColor, QIcon, QPainter, QPen, QPixmap, QPolygon, QPolygonF
from PySide6.QtWidgets import (QButtonGroup, QCheckBox, QColorDialog, QComboBox,
                               QDialog, QDialogButtonBox, QFrame, QHBoxLayout,
                               QInputDialog, QLabel, QLineEdit, QListWidget,
                               QListWidgetItem, QMenu, QMessageBox, QPlainTextEdit,
                               QPushButton, QTableWidget, QTableWidgetItem,
                               QToolButton, QVBoxLayout, QWidget)
from openpyxl.styles import Alignment, Font, PatternFill

from atlas_formula import col_to_num, num_to_col
from atlas_model import Book, RANGES_SHEET
from atlas_theme import LAYER_COLORS, qcolor

# The Biomes six ticks (also the preset of that name)
BLOCKS = [0, 4, 8, 12, 16, 20]
DAY = {8, 12, 16}            # 08:00-20:00
NIGHT = {20, 0, 4}           # 20:00-08:00
LEVELS = {3: "common", 2: "uncommon", 1: "rare"}
ALPHA = {3: 215, 2: 135, 1: 70}
BRUSH_OFF = 9                # kept for older callers
HEADERS = ["Layer", "Map", "Times", "Abundance", "Squares", "Colour", "Style"]
STYLES = {"area": "Area", "path": "Path", "marker": "Marker"}
FADE = 0.3                   # other layers' areas, when only the selected one is in focus


# ---------------------------------------------------------------- time and rarity


class TimeScheme:
    """How a workbook's world keeps time."""

    PRESETS = {
        "none": "No time (things are just there)",
        "daynight": "Day / Night",
        "four": "Dawn / Day / Dusk / Night",
        "hourly": "Hourly (24)",
        "biomes": "Six 4-hour ticks (00, 04 … 20), grouped Day / Night",
        "custom": "My own periods",
    }

    def __init__(self, preset, periods, groups):
        self.preset = preset
        self.periods = list(periods)             # [(id, label)]
        self.ids = [p for p, _ in self.periods]
        self.labels = dict(self.periods)
        self.groups = [(n, frozenset(g)) for n, g in groups if g]
        self.all = frozenset(self.ids)
        # the parts a square splits into in the "split" view (2 to 4), or none
        if 2 <= len(self.groups) <= 4 and frozenset().union(*[g for _, g in self.groups]) == self.all:
            self.split = list(self.groups)
        elif 2 <= len(self.ids) <= 4:
            self.split = [(self.labels[p], frozenset([p])) for p in self.ids]
        else:
            self.split = []

    @classmethod
    def from_meta(cls, meta):
        t = meta.get("time") or {"preset": "none"}
        preset = t.get("preset", "none")
        if preset == "daynight":
            return cls(preset, [("day", "Day"), ("night", "Night")], [])
        if preset == "four":
            return cls(preset, [("dawn", "Dawn"), ("day", "Day"), ("dusk", "Dusk"), ("night", "Night")], [])
        if preset == "hourly":
            return cls(preset, [(h, f"{h:02d}:00") for h in range(24)],
                       [("Day", set(range(6, 18))), ("Night", set(range(0, 6)) | set(range(18, 24)))])
        if preset == "biomes":
            return cls(preset, [(b, f"{b:02d}:00") for b in BLOCKS], [("Day", DAY), ("Night", NIGHT)])
        if preset == "custom":
            names = [n for n in t.get("periods", []) if str(n).strip()]
            if not names:
                return cls("none", [("always", "Always")], [])
            groups = [(g, set(m) & set(names)) for g, m in (t.get("groups") or {}).items()]
            return cls(preset, [(n, n) for n in names], groups)
        return cls("none", [("always", "Always")], [])

    def has_time(self):
        return len(self.ids) > 1

    def all_label(self):
        return "Always" if not self.has_time() else "All day"

    def label(self, pid):
        return self.labels.get(pid, str(pid))

    def times_label(self, ids):
        ids = frozenset(ids) & self.all
        if ids == self.all:
            return self.all_label()
        for name, g in self.groups:
            if ids == g:
                return name
        return ", ".join(self.label(p) for p in self.ids if p in ids)

    def parse(self, text):
        t = str(text or "").strip()
        low = t.lower()
        if low in ("", "all", "all day", "always", "any time", "anytime"):
            return set(self.all)
        for name, g in self.groups:
            if low == name.lower():
                return set(g)
        out = set()
        for part in re.split(r"[,;]+", t):
            part = part.strip()
            if not part:
                continue
            for pid, lab in self.periods:
                if part.lower() in (str(lab).lower(), str(pid).lower()):
                    out.add(pid)
                    break
            else:
                m = re.match(r"(\d{1,2})(?::00)?$", part)
                if m and int(m.group(1)) in self.all:
                    out.add(int(m.group(1)))
        if not out:
            # old files: "Day" / "Night" / tick lists written for the six ticks
            for name, g in (("day", DAY), ("night", NIGHT)):
                if low == name and set(g) <= set(self.all):
                    return set(g)
        return out or set(self.all)

    def options(self):
        """[(label, ids)] offered for painting and viewing: all, groups, each period."""
        out = [(self.all_label(), self.all)]
        out += [(n, g) for n, g in self.groups]
        out += [(self.label(p), frozenset([p])) for p in self.ids]
        return out if self.has_time() else out[:1]

    def tone(self, ids):
        """0 (day/full colour) to 1 (night shade), from the names involved."""
        names = " ".join(self.label(p).lower() for p in ids) + " " + " ".join(
            n.lower() for n, g in self.groups if g and g <= frozenset(ids))
        if "night" in names and not any(w in names for w in ("day", "dawn", "dusk", "morning")):
            return 1.0
        if any(w in names for w in ("dawn", "dusk", "twilight", "evening", "morning", "sunset", "sunrise")) \
                and "day" not in names.split():
            return 0.5
        if isinstance(next(iter(ids), None), int) and self.preset in ("biomes", "hourly"):
            night = next((g for n, g in self.groups if n.lower() == "night"), frozenset())
            if ids and frozenset(ids) <= night:
                return 1.0
        return 0.0


def scheme_of(book):
    """The workbook's time scheme. Workbooks that already have layers keep
    the Biomes ticks; new ones start with no time."""
    if "time" not in book.meta:
        book.meta["time"] = {"preset": "biomes" if _had_layers(book) else "none"}
    return TimeScheme.from_meta(book.meta)


def rarity_of(book):
    """{3: name, 2: name, 1: name} (fewer when fewer names), or None when off."""
    if "rarity" not in book.meta:
        book.meta["rarity"] = ["Common", "Uncommon", "Rare"] if _had_layers(book) else None
    names = book.meta.get("rarity")
    if not names:
        return None
    return {3 - i: n for i, n in enumerate(names[:3])}


def _had_layers(book):
    return bool(book.meta.get("ranges_sheet")) or book.sheet(RANGES_SHEET) is not None


def level_name(book, lv):
    r = rarity_of(book)
    return (r or {}).get(lv, "here") if lv else "not here"


def alpha_for(book, lv):
    return ALPHA[3] if rarity_of(book) is None else ALPHA.get(lv, ALPHA[3])


def times_label(blocks, book=None):
    """Kept for callers that only know the Biomes ticks."""
    scheme = scheme_of(book) if book is not None else TimeScheme.from_meta({"time": {"preset": "biomes"}})
    return scheme.times_label(blocks)


def block_label(b):
    return f"{b:02d}:00" if isinstance(b, int) else str(b)


def mix(hex_a, hex_b, t):
    a = [int(hex_a[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(hex_b[i:i + 2], 16) for i in (0, 2, 4)]
    return "%02X%02X%02X" % tuple(round(x * (1 - t) + y * t) for x, y in zip(a, b))


def night_shade(hex6):
    """The night version of a layer colour: darker and a little bluer."""
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, s, v = colorsys.rgb_to_hsv(r, g, b)
    r, g, b = colorsys.hsv_to_rgb(h, min(1.0, s * 1.1), v * 0.52)
    nr, ng, nb = 0.20, 0.24, 0.42
    m = 0.28
    r, g, b = r * (1 - m) + nr * m, g * (1 - m) + ng * m, b * (1 - m) + nb * m
    return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))


def toned(hex6, t):
    return hex6 if t <= 0 else night_shade(hex6) if t >= 1 else mix(hex6, night_shade(hex6), t)


# ---------------------------------------------------------------- squares


def square_label(sq):
    return f"{num_to_col(sq[0])}{sq[1]}"


_SQ = re.compile(r"([A-Za-z]{1,3})(\d+)")


def parse_squares(text):
    out = set()
    for part in re.split(r"[,;\s]+", str(text or "")):
        if not part:
            continue
        ends = part.split(":")
        a = _SQ.fullmatch(ends[0])
        b = _SQ.fullmatch(ends[-1])
        if not a or not b:
            continue
        x1, y1 = col_to_num(a.group(1)), int(a.group(2))
        x2, y2 = col_to_num(b.group(1)), int(b.group(2))
        for x in range(min(x1, x2), max(x1, x2) + 1):
            for y in range(min(y1, y2), max(y1, y2) + 1):
                out.add((x, y))
    return out


def squares_text(squares):
    """Compress a set of squares into rectangles: 'G9:K12, I13'."""
    left = set(squares)
    parts = []
    for sq in sorted(squares, key=lambda s: (s[1], s[0])):
        if sq not in left:
            continue
        x, y = sq
        x2 = x
        while (x2 + 1, y) in left:
            x2 += 1
        y2 = y
        while all((xx, y2 + 1) in left for xx in range(x, x2 + 1)):
            y2 += 1
        for xx in range(x, x2 + 1):
            for yy in range(y, y2 + 1):
                left.discard((xx, yy))
        parts.append(square_label((x, y)) if (x, y) == (x2, y2)
                     else f"{square_label((x, y))}:{square_label((x2, y2))}")
    return ", ".join(parts)


# ---------------------------------------------------------------- layers


class Layer:
    def __init__(self, name, color):
        self.name = name
        self.color = color
        self.visible = True
        self.style = "area"      # area (fill), path (a line through squares) or marker (a dot)
        self.maps = {}           # map title -> {period id: {(x, y): level}}

    def level(self, map_title, block, sq):
        return self.maps.get(map_title, {}).get(block, {}).get(sq, 0)

    def present(self, map_title, sq):
        return any(cells.get(sq) for cells in self.maps.get(map_title, {}).values())

    def has_map(self, map_title):
        return any(self.maps.get(map_title, {}).values())


class Ranges:
    def __init__(self, scheme=None):
        self.layers = []
        self.scheme = scheme or TimeScheme.from_meta({"time": {"preset": "biomes"}})

    def get(self, name):
        for layer in self.layers:
            if layer.name.lower() == name.lower():
                return layer
        return None

    def add(self, name, color=None):
        if color is None:
            color = LAYER_COLORS[len(self.layers) % len(LAYER_COLORS)]
        layer = Layer(name, color)
        self.layers.append(layer)
        return layer

    def paint(self, layer, map_title, blocks, squares, level):
        if layer.style != "area" or not self.scheme.has_time():
            blocks = self.scheme.all          # paths, markers and timeless worlds: just there
        if layer.style != "area":
            level = 3 if level else 0
        m = layer.maps.setdefault(map_title, {})
        for b in blocks:
            cells = m.setdefault(b, {})
            for sq in squares:
                if level:
                    cells[sq] = level
                else:
                    cells.pop(sq, None)

    def who(self, map_title, sq):
        """[(layer, {period: level})] for one square."""
        out = []
        for layer in self.layers:
            levels = {b: layer.level(map_title, b, sq) for b in self.scheme.ids}
            if any(levels.values()):
                out.append((layer, levels))
        return out

    def rename_map(self, old, new):
        for layer in self.layers:
            if old in layer.maps:
                layer.maps[new] = layer.maps.pop(old)

    def drop_map(self, title):
        for layer in self.layers:
            layer.maps.pop(title, None)

    def convert(self, new):
        """Move every layer to a new time scheme. Periods both schemes share
        keep their squares; anything else becomes 'all the time'."""
        old = self.scheme
        common = old.all & new.all
        for layer in self.layers:
            for m, periods in list(layer.maps.items()):
                union = {}
                for cells in periods.values():
                    for sq, lv in cells.items():
                        if lv:
                            union[sq] = max(union.get(sq, 0), lv)
                kept = {p: dict(c) for p, c in periods.items() if p in common}
                covered = {sq for c in kept.values() for sq in c}
                fresh = {p: dict(kept.get(p, {})) for p in new.ids}
                for sq, lv in union.items():
                    if sq not in covered:
                        for p in new.ids:
                            fresh[p][sq] = lv
                layer.maps[m] = fresh
        self.scheme = new

    # -- the Ranges sheet
    @classmethod
    def load(cls, book):
        scheme = scheme_of(book)
        rarity_of(book)
        r = cls(scheme)
        title = book.meta.get("ranges_sheet", RANGES_SHEET)
        ws = book.sheet(title)
        heads = [str(ws.cell(1, i + 1).value or "") for i in range(6)] if ws is not None else []
        if not heads or heads[0] not in ("Layer", "Creature") or heads[1:] != HEADERS[1:6]:
            return r
        book.meta["ranges_sheet"] = title
        names = {v.lower(): k for k, v in (rarity_of(book) or {}).items()}
        names.update({v: k for k, v in LEVELS.items()})
        for row in ws.iter_rows(min_row=2, values_only=True):
            row = list(row) + [None] * 7
            name, map_title, times, level, squares, color, style = row[:7]
            if not name:
                continue
            layer = r.get(str(name)) or r.add(str(name), _hex(color))
            if _hex(color):
                layer.color = _hex(color)
            if str(style or "").strip().lower() in STYLES:
                layer.style = str(style).strip().lower()
            if map_title:
                lv = names.get(str(level or "").strip().lower(), 3)
                r.paint(layer, str(map_title), scheme.parse(times), parse_squares(squares), lv)
        return r

    def write(self, book):
        title = book.meta.get("ranges_sheet")
        if title is None:
            if not self.layers:
                return
            title = RANGES_SHEET if book.sheet(RANGES_SHEET) is None else book.unique_title("Ranges (Atlas)")
            book.meta["ranges_sheet"] = title
        ws = book.sheet(title)
        if ws is None:
            ws = book.wb.create_sheet(title)
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
        head_font = Font(name="Arial", size=10, bold=True, color="FFECEFF4")
        head_fill = PatternFill("solid", fgColor="FF3B4252")
        for i, h in enumerate(HEADERS, 1):
            cell = ws.cell(1, i, h)
            cell.font, cell.fill = head_font, head_fill
        for i, w in enumerate([26, 22, 22, 12, 60, 10, 10], 1):
            ws.column_dimensions[num_to_col(i)].width = w
        ws.freeze_panes = "A2"
        scheme = self.scheme
        rarity = rarity_of(book)
        row = 2
        for layer in self.layers:
            wrote = False
            for map_title, periods in layer.maps.items():
                rows_out = []
                for lv in (3, 2, 1):
                    by_squares = {}
                    for p in scheme.ids:
                        sqs = frozenset(sq for sq, v in periods.get(p, {}).items() if v == lv)
                        if sqs:
                            by_squares.setdefault(sqs, set()).add(p)
                    for sqs, ps in by_squares.items():
                        rows_out.append((lv, ps, sqs))
                rows_out.sort(key=lambda x: (min(scheme.ids.index(p) for p in x[1]), -x[0]))
                for lv, ps, sqs in rows_out:
                    abundance = (rarity or {}).get(lv, "") if rarity else ""
                    values = [layer.name, map_title, scheme.times_label(ps), abundance, squares_text(sqs),
                              layer.color, layer.style]
                    for i, v in enumerate(values, 1):
                        ws.cell(row, i, v).alignment = Alignment(vertical="top", wrap_text=(i == 5))
                    ws.cell(row, 6).fill = PatternFill("solid", fgColor="FF" + layer.color)
                    row += 1
                    wrote = True
            if not wrote:
                ws.cell(row, 1, layer.name)
                ws.cell(row, 6, layer.color).fill = PatternFill("solid", fgColor="FF" + layer.color)
                ws.cell(row, 7, layer.style)
                row += 1


def _hex(v):
    s = str(v or "").strip().lstrip("#").upper()
    return s if re.fullmatch(r"[0-9A-F]{6}", s) else None


Book.ranges_writer = staticmethod(lambda book: book.ranges.write(book) if book.ranges is not None else None)


# ---------------------------------------------------------------- maps


def detect_map(ws):
    """Find a lettered row (A, B, C ...) with a numbered column (1, 2, 3 ...)
    down its left side. Returns {"row", "col", "rows", "cols"} for square A1,
    or None."""
    cells = ws._cells
    for (r, c), cell in sorted(cells.items()):
        if r > 40:
            break
        if str(cell.value).strip() != "A":
            continue
        n = 0
        while str((cells.get((r, c + n)) or _Empty).value).strip() == num_to_col(n + 1):
            n += 1
        if n < 4:
            continue
        m = 0
        while True:
            v = (cells.get((r + 1 + m, c - 1)) or _Empty).value
            try:
                ok = float(v) == m + 1
            except (TypeError, ValueError):
                ok = False
            if not ok:
                break
            m += 1
        if m >= 3:
            return {"row": r + 1, "col": c, "rows": m, "cols": n}
    return None


class _Empty:
    value = None


def map_square(info, row, col):
    """Map square (x, y) for a sheet cell, or None if outside the grid."""
    x, y = col - info["col"] + 1, row - info["row"] + 1
    if 1 <= x <= info["cols"] and 1 <= y <= info["rows"]:
        return (x, y)
    return None


def square_cell(info, sq):
    return info["row"] + sq[1] - 1, info["col"] + sq[0] - 1


# ---------------------------------------------------------------- view state


class ViewState:
    @property
    def split(self):
        return self.view == "split"

    @split.setter
    def split(self, on):
        self.view = "split" if on else "all"

    def __init__(self):
        self.blocks = None           # the times Add/brush paints (None = all of them)
        self.view = "split"          # "split", "all", or a frozenset of periods
        self.level = 3
        self.painting = False        # a tool is picked up (drag on the map)
        self.erasing = False         # ... and it's the eraser
        self.layer = None            # the selected layer
        self.solo = None
        self.all_areas = False       # False: other layers' areas fade (paths/markers always show)

    def paint_blocks(self, scheme):
        if self.blocks is None or not set(self.blocks) <= set(scheme.all) or not self.blocks:
            return set(scheme.all)
        return set(self.blocks)


def view_parts(scheme, view):
    """[(label, ids, tone)] the view shows: several parts in split view."""
    if view == "split" and scheme.split:
        n = len(scheme.split)
        return [(lab, ids, scheme.tone(ids)) for lab, ids in scheme.split]
    ids = scheme.all if view in ("all", "split", None) else frozenset(view) & scheme.all or scheme.all
    tone = scheme.tone(ids) if ids != scheme.all else 0.0
    return [(scheme.times_label(ids), ids, tone)]


def view_blocks(view, scheme=None):
    scheme = scheme or TimeScheme.from_meta({"time": {"preset": "biomes"}})
    named = {"day": DAY, "night": NIGHT, "all": scheme.all}
    if isinstance(view, str) and view in named:
        return set(named[view])
    if isinstance(view, int):
        return {view}
    return set(view) if view and view != "split" else set(scheme.all)


# ---------------------------------------------------------------- paths


def path_end(book, layer, map_title, sq, toward, info):
    """How a path ends in a square with one neighbour: 'run' (straight on to
    the far edge) or 'end' (stops in the square). Ends on the edge of the map
    run off it unless set otherwise."""
    mode = book.meta.get("path_ends", {}).get(layer.name, {}).get(map_title, {}).get(square_label(sq))
    if mode:
        return mode
    dx, dy = toward
    nx, ny = sq[0] - dx, sq[1] - dy
    off_map = not (1 <= nx <= info["cols"] and 1 <= ny <= info["rows"])
    return "run" if off_map else "end"


def path_ends_here(book, map_title, sq):
    """[(layer, toward)] for each path that ends in this square."""
    out = []
    for layer in book.ranges.layers:
        if layer.style != "path" or not layer.present(map_title, sq):
            continue
        dirs = []
        for dx in (-1, 0, 1):
            for dy in (-1, 0, 1):
                if (dx or dy) and layer.present(map_title, (sq[0] + dx, sq[1] + dy)):
                    if dx and dy and (layer.present(map_title, (sq[0] + dx, sq[1]))
                                      or layer.present(map_title, (sq[0], sq[1] + dy))):
                        continue
                    dirs.append((dx, dy))
        if len(dirs) == 1:
            out.append((layer, dirs[0]))
    return out


# ---------------------------------------------------------------- drawing


class Overlay:
    """Paints layers onto map squares (called by the grid). Areas fill (split
    by time in the split view; several areas side by side), paths draw a line
    joining neighbouring squares, markers draw a dot."""

    def __init__(self, window):
        self.window = window

    def _levels(self, layer, map_title, sq, parts):
        return [max((layer.level(map_title, b, sq) for b in ids), default=0) for _, ids, _ in parts]

    def line_through(self, ws, row, col):
        """True when a path or marker is drawn on this square (labels get a backing)."""
        book = self.window.book
        if book is None or book.ranges is None or not self.window.layers_showing():
            return False
        info = book.meta["maps"].get(ws.title)
        sq = map_square(info, row, col) if info else None
        if sq is None:
            return False
        st = self.window.range_state
        for layer in book.ranges.layers:
            if layer.style == "area" or (st.solo and layer.name != st.solo) or (not st.solo and not layer.visible):
                continue
            if layer.present(ws.title, sq):
                return True
        return False

    def paint(self, p, ws, row, col, rect):
        book = self.window.book
        if book is None or book.ranges is None or not self.window.layers_showing():
            return
        info = book.meta["maps"].get(ws.title)
        if not info:
            return
        sq = map_square(info, row, col)
        if sq is None:
            return
        st = self.window.range_state
        scheme = book.ranges.scheme
        parts = view_parts(scheme, st.view)
        areas, lines, marks = [], [], []
        for layer in book.ranges.layers:
            if st.solo and layer.name != st.solo:
                continue
            if not st.solo and not layer.visible:
                continue
            if layer.style == "area":
                lvs = self._levels(layer, ws.title, sq, parts)
                if any(lvs):
                    focus = st.solo or st.all_areas or layer.name == st.layer
                    areas.append((layer, lvs, 1.0 if focus else FADE))
            elif layer.present(ws.title, sq):
                (lines if layer.style == "path" else marks).append(layer)
        if not (areas or lines or marks):
            return
        inner = rect.adjusted(0, 0, -1, -1)
        p.save()
        p.setRenderHint(QPainter.Antialiasing, True)
        p.setPen(Qt.NoPen)
        k = len(areas)
        for i, (layer, lvs, fade) in enumerate(areas):
            x1 = inner.left() + inner.width() * i // k
            x2 = inner.left() + inner.width() * (i + 1) // k
            stripe = QRect(x1, inner.top(), x2 - x1, inner.height())
            shapes = _part_shapes(stripe, len(parts))
            for (lab, ids, tone), lv, shape in zip(parts, lvs, shapes):
                if lv:
                    p.setBrush(qcolor(toned(layer.color, tone), int(alpha_for(book, lv) * fade)))
                    p.drawPolygon(shape)
            if len(parts) > 1 and sum(1 for lv in lvs if lv) > 1 and fade == 1.0:
                p.setPen(QPen(QColor(20, 20, 26, 100), 1))
                for a, b in _part_seams(stripe, len(parts)):
                    p.drawLine(a, b)
                p.setPen(Qt.NoPen)
        c = QPointF(inner.center())
        size = min(inner.width(), inner.height())
        for layer in lines:
            dirs = []
            for dx in (-1, 0, 1):
                for dy in (-1, 0, 1):
                    if (dx, dy) == (0, 0) or not layer.present(ws.title, (sq[0] + dx, sq[1] + dy)):
                        continue
                    if dx and dy and (layer.present(ws.title, (sq[0] + dx, sq[1]))
                                      or layer.present(ws.title, (sq[0], sq[1] + dy))):
                        continue
                    dirs.append((dx, dy))
            if len(dirs) == 1 and path_end(book, layer, ws.title, sq, dirs[0], info) == "run":
                dx, dy = dirs[0]
                dirs.append((-dx, -dy))
            ends = [QPointF(c.x() + dx * inner.width() / 2, c.y() + dy * inner.height() / 2) for dx, dy in dirs]
            width = max(3.0, size * 0.17)
            for pen_color, extra in ((QColor(20, 20, 26, 110), 2.0), (qcolor(layer.color), 0.0)):
                w2 = width + extra
                p.setPen(QPen(pen_color, w2, Qt.SolidLine, Qt.FlatCap, Qt.RoundJoin))
                for e in ends:
                    p.drawLine(c, e)
                p.setPen(Qt.NoPen)
                p.setBrush(pen_color)
                p.drawEllipse(c, w2 / 2, w2 / 2)
                p.setBrush(Qt.NoBrush)
        for i, layer in enumerate(marks):
            r = size * 0.2
            off = (i - (len(marks) - 1) / 2) * r * 2.2
            p.setPen(QPen(QColor(20, 20, 26, 170), 2))
            p.setBrush(qcolor(layer.color))
            p.drawEllipse(QPointF(c.x() + off, c.y()), r, r)
        p.restore()


def _part_shapes(r, n):
    """Polygons a square splits into: whole, diagonal halves, or wedges."""
    tl, tr = QPointF(r.left(), r.top()), QPointF(r.right() + 1, r.top())
    bl, br = QPointF(r.left(), r.bottom() + 1), QPointF(r.right() + 1, r.bottom() + 1)
    c = QPointF(r.center()) + QPointF(0.5, 0.5)
    if n <= 1:
        return [QPolygonF([tl, tr, br, bl])]
    if n == 2:
        return [QPolygonF([tl, tr, bl]), QPolygonF([tr, br, bl])]
    top, right, bottom, left = (QPolygonF([tl, tr, c]), QPolygonF([tr, br, c]),
                                QPolygonF([br, bl, c]), QPolygonF([bl, tl, c]))
    if n == 3:
        return [top, right, QPolygonF([br, bl, tl, c])]
    return [top, right, bottom, left]


def _part_seams(r, n):
    tl, tr = QPointF(r.left(), r.top()), QPointF(r.right() + 1, r.top())
    bl, br = QPointF(r.left(), r.bottom() + 1), QPointF(r.right() + 1, r.bottom() + 1)
    if n == 2:
        return [(tr, bl)]
    return [(tl, br), (tr, bl)]


class Brush:
    """Paint by dragging (optional): the grid's tool while it's switched on."""

    def __init__(self, window):
        self.window = window
        self.stroke = None

    def active(self):
        w = self.window
        st = w.range_state
        return (st.painting and st.layer is not None and w.book is not None
                and w.current_ws() is not None and w.current_ws().title in w.book.meta["maps"])

    def _apply(self, row, col):
        w = self.window
        info = w.book.meta["maps"][w.current_ws().title]
        sq = map_square(info, row, col)
        if sq is None or sq in self.stroke["done"]:
            return
        self.stroke["done"].add(sq)
        layer = w.book.ranges.get(w.range_state.layer)
        if layer is None:
            return
        st = w.range_state
        w.book.ranges.paint(layer, w.current_ws().title, st.paint_blocks(w.book.ranges.scheme), [sq],
                            0 if self.stroke["erase"] else (st.level or 3))
        w.grid.viewport().update()

    def press(self, row, col, e):
        if not self.active() or e.button() not in (Qt.LeftButton,):
            return False
        info = self.window.book.meta["maps"][self.window.current_ws().title]
        if map_square(info, row, col) is None:
            return False
        if self.stroke is not None:
            self.window.book.done("ranges")
        self.window.book.begin([])
        erase = bool(e.modifiers() & Qt.ShiftModifier) or self.window.range_state.erasing
        self.stroke = {"done": set(), "erase": erase}
        self._apply(row, col)
        return True

    def move(self, row, col, e):
        if self.stroke is not None:
            self._apply(row, col)

    def release(self, e):
        if self.stroke is not None:
            self.stroke = None
            self.window.book.done("ranges")

    def hover(self, row, col):
        pass

    def stop(self):
        self.window.stop_painting()


# ---------------------------------------------------------------- the panel


def tool_icon(kind, fill_hex, line_hex, size=26):
    """Brush (its tip in the layer's colour) or eraser, drawn in code."""
    pm = QPixmap(size, size)
    pm.fill(Qt.transparent)
    p = QPainter(pm)
    p.setRenderHint(QPainter.Antialiasing)
    p.scale(size / 26, size / 26)
    line = qcolor(line_hex)
    if kind == "brush":
        p.setPen(QPen(line, 2.2, Qt.SolidLine, Qt.RoundCap))
        p.drawLine(QPointF(20.5, 4.5), QPointF(12.5, 12.5))           # handle
        p.setPen(QPen(line, 1.4))
        p.setBrush(qcolor(fill_hex))
        p.drawPolygon(QPolygonF([QPointF(12.0, 11.0), QPointF(15.0, 14.0),        # bristles
                                 QPointF(9.5, 21.5), QPointF(4.0, 22.0), QPointF(4.5, 16.5)]))
    else:
        p.translate(13, 13)
        p.rotate(-40)
        p.setPen(QPen(line, 1.4))
        p.setBrush(qcolor("BF616A"))
        p.drawRoundedRect(QRectF(-9, -5, 9, 10), 2, 2)
        p.setBrush(qcolor("ECEFF4"))
        p.drawRoundedRect(QRectF(0, -5, 9, 10), 2, 2)
    p.end()
    return QIcon(pm)


def swatch(hex6, size=14):
    pm = QPixmap(size, size)
    pm.fill(Qt.transparent)
    p = QPainter(pm)
    p.setRenderHint(QPainter.Antialiasing)
    p.setPen(Qt.NoPen)
    p.setBrush(qcolor(hex6))
    p.drawRoundedRect(0, 0, size, size, 3, 3)
    p.end()
    return QIcon(pm)


class RangesPanel(QWidget):
    """The Layers section of the side panel."""

    changed = Signal()

    def __init__(self, window):
        super().__init__()
        self.window = window
        self.st = window.range_state
        lay = QVBoxLayout(self)
        lay.setContentsMargins(10, 8, 10, 10)
        lay.setSpacing(8)

        self.title = QLabel()
        self.title.setObjectName("heading")
        self.title.setWordWrap(True)
        lay.addWidget(self.title)

        self.no_map = QWidget()
        nm = QVBoxLayout(self.no_map)
        nm.setContentsMargins(0, 0, 0, 0)
        self.no_map_text = QLabel()
        self.no_map_text.setObjectName("muted")
        self.no_map_text.setWordWrap(True)
        nm.addWidget(self.no_map_text)
        self.btn_detect = QPushButton("Find the map grid on this sheet")
        self.btn_detect.clicked.connect(self.detect)
        self.btn_use_sel = QPushButton("Use the selected squares as the map")
        self.btn_use_sel.clicked.connect(self.use_selection)
        nm.addWidget(self.btn_detect)
        nm.addWidget(self.btn_use_sel)
        lay.addWidget(self.no_map)

        self.body = QWidget()
        b = QVBoxLayout(self.body)
        b.setContentsMargins(0, 0, 0, 0)
        b.setSpacing(8)

        # the layers
        self.list = QListWidget()
        self.list.setToolTip("Anything with a footprint on the map: creatures, plants, characters, paths, quests.\n"
                             "Select one to work on it. Tick to show; double-click to show only that one.")
        self.list.itemChanged.connect(self.item_changed)
        self.list.currentItemChanged.connect(self.current_layer_changed)
        self.list.itemDoubleClicked.connect(self.solo_toggle)
        self.list.setContextMenuPolicy(Qt.CustomContextMenu)
        self.list.customContextMenuRequested.connect(self.layer_menu)
        b.addWidget(self.list)
        self.empty_note = QLabel("No layers yet. Click New, then select squares on the map and click Add.")
        self.empty_note.setObjectName("muted")
        self.empty_note.setWordWrap(True)
        b.addWidget(self.empty_note)
        row = QHBoxLayout()
        for label, fn in (("New", self.new_layer), ("Rename", self.rename_layer),
                          ("Colour", self.recolor_layer), ("Delete", self.delete_layer)):
            btn = QPushButton(label)
            btn.clicked.connect(fn)
            row.addWidget(btn)
        b.addLayout(row)

        # ---- the work box: put squares in the selected layer
        self.work = QFrame()
        self.work.setObjectName("workbox")
        wk = QVBoxLayout(self.work)
        wk.setContentsMargins(10, 8, 10, 10)
        wk.setSpacing(7)
        self.work_title = QLabel()
        self.work_title.setObjectName("boxTitle")
        wk.addWidget(self.work_title)
        # options: when, and how common (each only if the workbook uses them)
        self.opts = QWidget()
        o = QHBoxLayout(self.opts)
        o.setContentsMargins(0, 0, 0, 0)
        self.for_label = QLabel("For")
        self.for_label.setObjectName("faint")
        self.for_box = QComboBox()
        self.for_box.activated.connect(self.for_chosen)
        self.as_label = QLabel("As")
        self.as_label.setObjectName("faint")
        self.as_box = QComboBox()
        self.as_box.activated.connect(self.as_chosen)
        for w_ in (self.for_label, self.for_box, self.as_label, self.as_box):
            o.addWidget(w_)
        o.addStretch()
        wk.addWidget(self.opts)
        # the tools: a brush that adds, an eraser that removes
        row = QHBoxLayout()
        row.setSpacing(6)
        self.paint_btn = QToolButton()
        self.paint_btn.setObjectName("toolBtn")
        self.paint_btn.setCheckable(True)
        self.paint_btn.setIconSize(QSize(26, 26))
        self.paint_btn.setToolTip("Brush: drag over squares to add them to this layer.\n"
                                  "Squares already selected are added as soon as you pick it. Esc puts it down.")
        self.paint_btn.toggled.connect(self.paint_toggled)
        self.erase_btn = QToolButton()
        self.erase_btn.setObjectName("toolBtn")
        self.erase_btn.setCheckable(True)
        self.erase_btn.setIconSize(QSize(26, 26))
        self.erase_btn.setToolTip("Eraser: drag over squares to take them out of this layer.\n"
                                  "Squares already selected are removed as soon as you pick it. Esc puts it down.")
        self.erase_btn.toggled.connect(self.erase_toggled)
        row.addWidget(self.paint_btn)
        row.addWidget(self.erase_btn)
        self.tool_hint = QLabel()
        self.tool_hint.setObjectName("muted")
        self.tool_hint.setWordWrap(True)
        row.addWidget(self.tool_hint, 1)
        wk.addLayout(row)
        b.addWidget(self.work)

        # ---- view
        lbl = QLabel("View")
        lbl.setObjectName("subhead")
        b.addWidget(lbl)
        self.view_row = QWidget()
        v = QHBoxLayout(self.view_row)
        v.setContentsMargins(0, 0, 0, 0)
        lbl = QLabel("Show")
        lbl.setObjectName("faint")
        v.addWidget(lbl)
        self.view_box = QComboBox()
        self.view_box.activated.connect(self.view_chosen)
        v.addWidget(self.view_box, 1)
        b.addWidget(self.view_row)
        row = QHBoxLayout()
        lbl = QLabel("Draw as")
        lbl.setObjectName("faint")
        row.addWidget(lbl)
        self.style_box = QComboBox()
        for key in STYLES:
            self.style_box.addItem({"area": "Area (fills squares)", "path": "Path (roads, tracks, rivers)",
                                    "marker": "Marker (one-off things)"}[key], key)
        self.style_box.activated.connect(self.style_chosen)
        row.addWidget(self.style_box, 1)
        b.addLayout(row)
        self.all_box = QCheckBox("Show every ticked area at full strength")
        self.all_box.setToolTip("Off: the selected layer's area is bright and the others fade.\n"
                                "Paths and markers always show.")
        self.all_box.toggled.connect(self.all_toggled)
        b.addWidget(self.all_box)
        self.solo_note = QPushButton()
        self.solo_note.setFlat(True)
        self.solo_note.setStyleSheet("text-align: left;")
        self.solo_note.clicked.connect(lambda: self.show_all())
        b.addWidget(self.solo_note)
        self.settings_btn = QPushButton("Times and rarity…")
        self.settings_btn.setFlat(True)
        self.settings_btn.setStyleSheet("text-align: left;")
        self.settings_btn.setToolTip("How this workbook's world keeps time, and whether layers have rarity")
        self.settings_btn.clicked.connect(lambda: self.window.edit_layer_settings())
        b.addWidget(self.settings_btn)

        # the selected square
        self.here_title = QLabel()
        self.here_title.setWordWrap(True)
        self.here_title.setTextFormat(Qt.RichText)
        b.addWidget(self.here_title)
        self.here_table = QTableWidget(0, 2)
        self.here_table.verticalHeader().hide()
        self.here_table.setShowGrid(True)
        self.here_table.setSelectionMode(QTableWidget.NoSelection)
        self.here_table.setFocusPolicy(Qt.NoFocus)
        self.here_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.here_table.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.here_table.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.here_table.horizontalHeader().setStretchLastSection(False)
        self.here_table.verticalHeader().setDefaultSectionSize(24)
        self.here_table.cellClicked.connect(self.here_clicked)
        b.addWidget(self.here_table)
        self.here_legend = QLabel()
        self.here_legend.setTextFormat(Qt.RichText)
        self.here_legend.setObjectName("faint")
        b.addWidget(self.here_legend)
        self.here_add = QComboBox()
        self.here_add.activated.connect(self.here_add_chosen)
        b.addWidget(self.here_add)
        self.here_sq = None
        self.here_layers = []
        self.here_cols = []
        b.addStretch(1)
        lay.addWidget(self.body, 1)
        lay.addStretch()
        self.refresh()

    # -- helpers
    def book(self):
        return self.window.book

    def scheme(self):
        return self.window.book.ranges.scheme

    def timeless_selected(self):
        layer = self.book().ranges.get(self.st.layer) if (self.book() and self.st.layer) else None
        return layer is not None and layer.style != "area"

    # -- state <-> controls
    def refresh(self):
        w = self.window
        ws = w.current_ws()
        book = w.book
        is_map = bool(book is not None and ws is not None and ws.title in book.meta["maps"])
        self.no_map.setVisible(not is_map)
        self.body.setVisible(is_map)
        if ws is None:
            self.title.setText("")
            self.no_map_text.setText("Open a workbook to work with layers.")
            self.btn_detect.hide()
            self.btn_use_sel.hide()
            return
        if not is_map:
            self.title.setText("")
            self.no_map_text.setText(f"{ws.title} isn't a map. Layers live on map sheets; "
                                     "if this sheet has a grid, Atlas can use it.")
            self.btn_detect.show()
            self.btn_use_sel.show()
            return
        self.title.setText(ws.title)
        self.all_box.blockSignals(True)
        self.all_box.setChecked(self.st.all_areas)
        self.all_box.blockSignals(False)
        for btn, on in ((self.paint_btn, self.st.painting and not self.st.erasing),
                        (self.erase_btn, self.st.painting and self.st.erasing)):
            btn.blockSignals(True)
            btn.setChecked(on)
            btn.blockSignals(False)
        self.fill_list()
        self.fill_options()
        self.update_notes()
        self.update_here()
        self.window.update_paint_bar()

    def fill_options(self):
        """The For / As / Show boxes, built from the workbook's own scheme."""
        scheme = self.scheme()
        st = self.st
        timeless = self.timeless_selected()
        # For
        self.for_box.blockSignals(True)
        self.for_box.clear()
        for label, ids in scheme.options():
            self.for_box.addItem(label, ids)
        want = frozenset(st.paint_blocks(scheme))
        for i in range(self.for_box.count()):
            if self.for_box.itemData(i) == want:
                self.for_box.setCurrentIndex(i)
                break
        self.for_box.blockSignals(False)
        show_for = scheme.has_time() and not timeless
        self.for_label.setVisible(show_for)
        self.for_box.setVisible(show_for)
        # As
        rarity = rarity_of(self.book())
        self.as_box.blockSignals(True)
        self.as_box.clear()
        for lv, name in sorted((rarity or {}).items(), reverse=True):
            self.as_box.addItem(name, lv)
        idx = self.as_box.findData(st.level)
        self.as_box.setCurrentIndex(max(0, idx))
        if rarity and idx < 0:
            st.level = 3
        self.as_box.blockSignals(False)
        show_as = bool(rarity) and not timeless
        self.as_label.setVisible(show_as)
        self.as_box.setVisible(show_as)
        self.opts.setVisible(show_for or show_as)
        # Show
        self.view_box.blockSignals(True)
        self.view_box.clear()
        if scheme.split:
            self.view_box.addItem(" / ".join(lab for lab, _ in scheme.split) + " (split squares)", "split")
        for label, ids in scheme.options():
            self.view_box.addItem(label, "all" if ids == scheme.all else ids)
        cur = st.view if st.view in ("split", "all") else frozenset(st.view)
        if cur == "split" and not scheme.split:
            st.view = cur = "all"
        for i in range(self.view_box.count()):
            if self.view_box.itemData(i) == cur:
                self.view_box.setCurrentIndex(i)
                break
        else:
            st.view = self.view_box.itemData(0)
            self.view_box.setCurrentIndex(0)
        self.view_box.blockSignals(False)
        self.view_row.setVisible(scheme.has_time())

    def fill_list(self):
        book = self.book()
        ws = self.window.current_ws()
        self.list.blockSignals(True)
        self.list.clear()
        current = None
        here = [l for l in book.ranges.layers if l.has_map(ws.title)]
        rest = [l for l in book.ranges.layers if not l.has_map(ws.title)]
        maps = [t for t in book.meta["maps"] if book.sheet(t) is not None]
        for group, layers in (("", here), ("Not on this map", rest)):
            if not layers:
                continue
            if group:
                head = QListWidgetItem(group)
                head.setFlags(Qt.NoItemFlags)
                f = head.font()
                f.setPointSizeF(max(7.0, f.pointSizeF() - 1))
                head.setFont(f)
                self.list.addItem(head)
            for layer in layers:
                text = layer.name
                if group:
                    where = [t for t in maps if layer.has_map(t)]
                    text += f"  · on {', '.join(where)}" if where else "  · nothing added yet"
                it = QListWidgetItem(swatch(layer.color), text)
                it.setData(Qt.UserRole, layer.name)
                it.setFlags(it.flags() | Qt.ItemIsUserCheckable)
                it.setCheckState(Qt.Checked if layer.visible else Qt.Unchecked)
                if group:
                    it.setForeground(QColor(self.window.theme["faint"]))
                self.list.addItem(it)
                if layer.name == self.st.layer:
                    current = it
        if current is None:
            for i in range(self.list.count()):
                if self.list.item(i).data(Qt.UserRole):
                    current = self.list.item(i)
                    self.st.layer = current.data(Qt.UserRole)
                    break
        if current is not None:
            self.list.setCurrentItem(current)
        self.list.blockSignals(False)
        rows = max(2, self.list.count())
        self.list.setFixedHeight(min(170, rows * self.list.sizeHintForRow(0 if self.list.count() else -1) + 8)
                                 if self.list.count() else 60)
        self.sync_style_box()

    def update_notes(self):
        st = self.st
        has = bool(self.book() is not None and self.book().ranges.layers)
        self.empty_note.setVisible(not has)
        layer = self.book().ranges.get(st.layer) if (has and st.layer) else None
        self.paint_btn.setEnabled(layer is not None)
        self.erase_btn.setEnabled(layer is not None)
        col = self.window.theme["text"].lstrip("#")
        self.paint_btn.setIcon(tool_icon("brush", layer.color if layer else col, col))
        self.erase_btn.setIcon(tool_icon("eraser", col, col))
        self.tool_hint.setText("Drag on the map to erase" if (st.painting and st.erasing) else
                               "Drag on the map to paint" if st.painting else
                               "Pick the brush, then drag on the map")
        if layer is None:
            self.work_title.setText("Pick or make a layer to put squares in")
        else:
            full = f"Put squares in {layer.name}"
            fm = self.work_title.fontMetrics()
            self.work_title.setText(fm.elidedText(full, Qt.ElideRight, 250))
            self.work_title.setToolTip(full)
        self.solo_note.setText(f"Only {st.solo} is showing · show all")
        self.solo_note.setVisible(bool(st.solo))

    def show_all(self):
        self.st.solo = None
        self.update_notes()
        self.window.grid.viewport().update()

    # -- the selected square
    def _here_columns(self):
        """[(header, ids)]: one column per period when there are few, per group
        when there are many, or a single 'Here' column without time."""
        scheme = self.scheme()
        if not scheme.has_time():
            return [("Here", scheme.all)]
        if len(scheme.ids) <= 8:
            cols = [(scheme.label(p)[:5].replace(":00", ""), frozenset([p])) for p in scheme.ids]
        elif scheme.groups:
            cols = [(n, g) for n, g in scheme.groups]
        else:
            cols = []
        return cols + [("All", scheme.all)]

    def update_here(self):
        w = self.window
        ws = w.current_ws()
        if ws is None or w.book is None or ws.title not in w.book.meta["maps"]:
            return
        info = w.book.meta["maps"][ws.title]
        sq = map_square(info, *w.grid.cur)
        self.here_sq = (ws.title, sq) if sq else None
        t = w.theme
        tbl = self.here_table
        tbl.blockSignals(True)
        if sq is None:
            self.here_title.setText("<span style='color:gray'>Select squares on the map to add them to a layer, "
                                    "or one square to see what's in it.</span>")
            tbl.setRowCount(0)
            tbl.hide()
            self.here_add.hide()
            self.here_legend.hide()
            tbl.blockSignals(False)
            return
        scheme = self.scheme()
        rarity = rarity_of(w.book)
        rows = w.book.ranges.who(ws.title, sq)
        self.here_title.setText(f"<span style='color:{t['group']}'><b>Square {square_label(sq)}</b></span>"
                                + (f" <span style='color:{t['muted']}'>· click to change</span>" if rows else
                                   f" <span style='color:{t['muted']}'>· nothing here yet</span>"))
        cols = self._here_columns()
        self.here_cols = cols
        tbl.setColumnCount(1 + len(cols))
        tbl.setHorizontalHeaderLabels([""] + [h for h, _ in cols])
        tbl.setRowCount(len(rows))
        self.here_layers = [layer.name for layer, _ in rows]
        for i, (layer, levels) in enumerate(rows):
            name = QTableWidgetItem(swatch(layer.color), layer.name + ("" if layer.style == "area" else f" ({layer.style})"))
            name.setFlags(Qt.ItemIsEnabled)
            name.setToolTip(layer.name)
            tbl.setItem(i, 0, name)
            for j, (head, ids) in enumerate(cols):
                lv = max((levels.get(p, 0) for p in ids), default=0)
                last = j == len(cols) - 1 and len(cols) > 1
                it = QTableWidgetItem("all" if last else ("" if lv else "–"))
                it.setTextAlignment(Qt.AlignCenter)
                it.setForeground(QColor(t["faint"] if not last else t["muted"]))
                if lv and not last:
                    tone = scheme.tone(ids) if layer.style == "area" else 0.0
                    it.setBackground(qcolor(toned(layer.color, tone), alpha_for(w.book, lv)))
                it.setFlags(Qt.ItemIsEnabled)
                what = level_name(w.book, lv)
                it.setToolTip(f"{layer.name}{'' if not scheme.has_time() else ' · ' + scheme.times_label(ids)}: "
                              f"{what}. Click to change.")
                tbl.setItem(i, 1 + j, it)
        tbl.resizeColumnToContents(0)
        tbl.setColumnWidth(0, min(120, max(80, tbl.columnWidth(0))))
        for j in range(len(cols)):
            tbl.setColumnWidth(1 + j, 34 if len(cols) > 1 else 60)
        tbl.setFixedHeight(tbl.horizontalHeader().height() + sum(tbl.rowHeight(i) for i in range(len(rows))) + 4)
        tbl.setVisible(bool(rows))
        tbl.setStyleSheet(f"QTableWidget {{ gridline-color: {t['panel']}; background: {t['window']}; }}")
        if rarity and any(layer.style == "area" for layer, _ in rows):
            def chip(a):
                c = qcolor(t["accent"].lstrip("#"), a)
                return (f"<span style='background-color: rgba({c.red()},{c.green()},{c.blue()},{a / 255:.2f})'>"
                        "&nbsp;&nbsp;&nbsp;&nbsp;</span>")
            self.here_legend.setText(" &nbsp; ".join(f"{chip(ALPHA[lv])} {name}" for lv, name in
                                                     sorted(rarity.items(), reverse=True)))
            self.here_legend.show()
        else:
            self.here_legend.hide()
        others = [l.name for l in w.book.ranges.layers if l.name not in self.here_layers]
        self.here_add.blockSignals(True)
        self.here_add.clear()
        self.here_add.addItem("Add a layer to this square…", None)
        for n in others:
            self.here_add.addItem(swatch(w.book.ranges.get(n).color), n, n)
        self.here_add.addItem("New layer…", "__new__")
        self.here_add.blockSignals(False)
        self.here_add.show()
        tbl.blockSignals(False)

    def here_clicked(self, row, col):
        if self.here_sq is None or col == 0 or row >= len(self.here_layers):
            return
        w = self.window
        map_title, sq = self.here_sq
        layer = w.book.ranges.get(self.here_layers[row])
        if layer is None:
            return
        head, ids = self.here_cols[col - 1]
        rarity = rarity_of(w.book)
        cur = max((layer.level(map_title, p, sq) for p in ids), default=0)
        if layer.style != "area" or not rarity:
            level = 0 if cur else 3                          # here / not here
        else:
            order = sorted(rarity, reverse=True)             # e.g. [3, 2, 1]
            seq = order + [0]
            level = seq[(seq.index(cur) + 1) % len(seq)] if cur in seq else order[0]
        w.book.begin([])
        w.book.ranges.paint(layer, map_title, ids, [sq], level)
        w.book.done("ranges")

    def here_add_chosen(self, idx):
        data = self.here_add.itemData(idx)
        if not data or self.here_sq is None:
            return
        w = self.window
        map_title, sq = self.here_sq
        if data == "__new__":
            layer = self.new_layer()
            if layer is None:
                self.update_here()
                return
        else:
            layer = w.book.ranges.get(data)
        w.book.begin([])
        w.book.ranges.paint(layer, map_title, self.st.paint_blocks(self.scheme()), [sq], self.st.level or 3)
        w.book.done("ranges")

    # -- handlers
    def for_chosen(self, idx):
        self.set_blocks(self.for_box.itemData(idx))

    def as_chosen(self, idx):
        self.st.level = self.as_box.itemData(idx) or 3
        self.window.update_paint_bar()

    def view_chosen(self, idx):
        self.set_view(self.view_box.itemData(idx))

    def set_view(self, view):
        scheme = self.scheme()
        if view in ("day", "night") or isinstance(view, int):
            view = frozenset(view_blocks(view, scheme)) & scheme.all
        self.st.view = view
        self.fill_options()
        self.window.grid.viewport().update()

    def set_blocks(self, blocks, keep_split=False):
        """Which times Add and the brush put squares in. Doesn't touch the view."""
        self.st.blocks = set(blocks) if blocks else None
        self.fill_options()
        self.window.update_paint_bar()

    def sync_times(self):
        self.fill_options()

    def item_changed(self, it):
        layer = self.book().ranges.get(it.data(Qt.UserRole))
        if layer is not None:
            layer.visible = it.checkState() == Qt.Checked
            self.window.grid.viewport().update()

    def all_toggled(self, on):
        self.st.all_areas = on
        self.window.grid.viewport().update()

    def style_chosen(self, idx):
        layer = self._selected_layer()
        style = self.style_box.itemData(idx)
        if layer is None or style == layer.style:
            return
        self.book().begin([])
        live = self.book().ranges.get(layer.name)
        live.style = style
        if style != "area":
            for m in live.maps.values():
                where = set()
                for cells in m.values():
                    where |= {sq for sq, lv in cells.items() if lv}
                for p in self.scheme().ids:
                    m[p] = {sq: 3 for sq in where}
        self.book().done("ranges")

    def sync_style_box(self):
        layer = self._selected_layer() if self.list.currentItem() is not None else None
        self.style_box.setEnabled(layer is not None)
        if layer is not None:
            self.style_box.setCurrentIndex(self.style_box.findData(layer.style))

    def current_layer_changed(self, cur, prev):
        if cur is not None and cur.data(Qt.UserRole):
            self.st.layer = cur.data(Qt.UserRole)
            self.sync_style_box()
            self.fill_options()
            self.update_notes()
            self.window.grid.viewport().update()
            self.window.update_paint_bar()

    def solo_toggle(self, it):
        name = it.data(Qt.UserRole)
        if not name:
            return
        self.st.solo = None if self.st.solo == name else name
        self.update_notes()
        self.window.grid.viewport().update()

    def paint_toggled(self, on):
        self._tool(on, erasing=False)

    def erase_toggled(self, on):
        self._tool(on, erasing=True)

    def _tool(self, on, erasing):
        """Pick up (or put down) the brush or the eraser. Picking one up with
        several squares selected applies it to them straight away."""
        other = self.paint_btn if erasing else self.erase_btn
        if on:
            other.blockSignals(True)
            other.setChecked(False)
            other.blockSignals(False)
            if self.st.layer is None and self.new_layer() is None:
                (self.erase_btn if erasing else self.paint_btn).setChecked(False)
                return
        self.st.painting = on
        self.st.erasing = erasing and on
        if on and len(self.window.selected_squares()) > 1:
            self.window.add_selection_to_layer(remove=erasing)
        self.update_notes()
        self.window.update_paint_bar()

    def layer_menu(self, pos):
        it = self.list.itemAt(pos)
        menu = QMenu(self)
        if it is not None and it.data(Qt.UserRole):
            name = it.data(Qt.UserRole)
            menu.addAction(f"Show only {name}", lambda: self.solo_toggle(it))
            menu.addAction(f"Find “{name}” in the sheets", lambda: self.window.find_text(name))
            menu.addSeparator()
            menu.addAction("Rename…", self.rename_layer)
            menu.addAction("Colour…", self.recolor_layer)
            menu.addAction("Delete…", self.delete_layer)
        menu.addAction("New layer…", self.new_layer)
        menu.exec(self.list.mapToGlobal(pos))

    def _selected_layer(self):
        it = self.list.currentItem()
        name = it.data(Qt.UserRole) if it is not None else None
        return self.book().ranges.get(name) if name else None

    def new_layer(self, name=None):
        book = self.book()
        if name is None:
            name, ok = QInputDialog.getText(self, "New layer", "What's on the map? A creature, plant, character, "
                                                          "path, quest… (a name that matches a row elsewhere links to it)")
            if not ok:
                if self.st.painting and self.st.layer is None:
                    self.paint_btn.setChecked(False)
                return None
        name = name.strip()
        if not name:
            return None
        existing = book.ranges.get(name)
        if existing is not None:
            self.st.layer = existing.name
            self.refresh()
            return existing
        book.begin([])
        layer = book.ranges.add(name)
        book.done("ranges")
        self.st.layer = layer.name
        self.refresh()
        return layer

    def rename_layer(self):
        layer = self._selected_layer()
        if layer is None:
            return
        name, ok = QInputDialog.getText(self, "Rename layer", "New name:", text=layer.name)
        name = name.strip()
        if not ok or not name or name == layer.name:
            return
        if self.book().ranges.get(name) is not None and name.lower() != layer.name.lower():
            QMessageBox.information(self, "Rename layer", f"There's already a layer called {name}.")
            return
        self.book().begin([])
        old = layer.name
        self.book().ranges.get(old).name = name
        ends = self.book().meta.get("path_ends", {})
        if old in ends:
            ends[name] = ends.pop(old)
        if self.st.layer == old:
            self.st.layer = name
        if self.st.solo == old:
            self.st.solo = name
        self.book().done("ranges")

    def recolor_layer(self):
        layer = self._selected_layer()
        if layer is None:
            return
        hex6 = self.window.pick_color(layer.color, f"Colour for {layer.name}")
        if hex6 is None:
            return
        self.book().begin([])
        self.book().ranges.get(layer.name).color = hex6
        self.book().done("ranges")

    def delete_layer(self):
        layer = self._selected_layer()
        if layer is None:
            return
        book = self.book()
        here = self.window.current_ws().title
        maps = [t for t in book.meta["maps"] if book.sheet(t) is not None]
        counts = {t: len({sq for cells in layer.maps.get(t, {}).values() for sq, lv in cells.items() if lv})
                  for t in maps}
        painted = [t for t in maps if counts[t]]
        if painted:
            box = QMessageBox(self)
            box.setWindowTitle(layer.name)
            where = "; ".join(f"{t} ({counts[t]} square{'s' if counts[t] != 1 else ''})" for t in painted)
            box.setText(f"{layer.name} is on: {where}.")
            box.setInformativeText("Ctrl+Z undoes either choice.")
            only_here = box.addButton(f"Remove from {here}", QMessageBox.AcceptRole) \
                if here in painted and len(painted) > 1 else None
            everywhere = box.addButton("Delete the layer" if len(painted) == 1 else "Delete it everywhere",
                                       QMessageBox.DestructiveRole)
            box.addButton(QMessageBox.Cancel)
            box.exec()
            choice = box.clickedButton()
            if choice is only_here and only_here is not None:
                book.begin([])
                book.ranges.get(layer.name).maps.pop(here, None)
                book.done("ranges")
                return
            if choice is not everywhere:
                return
        book.begin([])
        book.ranges.layers = [l for l in book.ranges.layers if l.name != layer.name]
        book.meta.get("path_ends", {}).pop(layer.name, None)
        if self.st.layer == layer.name:
            self.st.layer = None
        if self.st.solo == layer.name:
            self.st.solo = None
        book.done("ranges")

    def detect(self):
        w = self.window
        ws = w.current_ws()
        info = detect_map(ws)
        if info is None:
            QMessageBox.information(self, "Find the map grid",
                                    "I couldn't find a lettered row (A, B, C…) with numbers (1, 2, 3…) down "
                                    "its left side. Select the squares of the map and use the other button.")
            return
        w.book.begin([])
        w.book.meta["maps"][ws.title] = info
        w.book.done("maps")

    def use_selection(self):
        w = self.window
        r1, c1, r2, c2 = w.grid.selection()
        if (r1, c1) == (r2, c2):
            QMessageBox.information(self, "Use selection as map",
                                    "Select all the squares of the map first (the top-left one becomes A1).")
            return
        w.book.begin([])
        w.book.meta["maps"][w.current_ws().title] = {"row": r1, "col": c1, "rows": r2 - r1 + 1, "cols": c2 - c1 + 1}
        w.book.done("maps")


# ---------------------------------------------------------------- settings dialog


class LayerSettingsDialog(QDialog):
    """Format -> Layer times and rarity."""

    def __init__(self, window):
        super().__init__(window)
        self.window = window
        book = window.book
        self.setWindowTitle("Layer times and rarity")
        self.resize(460, 420)
        lay = QVBoxLayout(self)
        t = book.meta.get("time") or {"preset": "none"}
        lbl = QLabel("How does this world keep time?")
        lbl.setObjectName("heading")
        lay.addWidget(lbl)
        self.preset = QComboBox()
        for key, label in TimeScheme.PRESETS.items():
            self.preset.addItem(label, key)
        self.preset.setCurrentIndex(max(0, self.preset.findData(t.get("preset", "none"))))
        lay.addWidget(self.preset)
        self.custom = QWidget()
        c = QVBoxLayout(self.custom)
        c.setContentsMargins(0, 0, 0, 0)
        hint = QLabel("One period per line, in order:")
        hint.setObjectName("muted")
        c.addWidget(hint)
        self.periods = QPlainTextEdit("\n".join(map(str, t.get("periods", ["Morning", "Afternoon", "Evening", "Night"]))))
        self.periods.setMaximumHeight(110)
        c.addWidget(self.periods)
        hint = QLabel("Groups (optional), one per line, e.g.  Day: Morning, Afternoon")
        hint.setObjectName("muted")
        c.addWidget(hint)
        self.groups = QPlainTextEdit("\n".join(f"{g}: {', '.join(map(str, m))}"
                                               for g, m in (t.get("groups") or {}).items()))
        self.groups.setMaximumHeight(70)
        c.addWidget(self.groups)
        lay.addWidget(self.custom)
        self.preset.currentIndexChanged.connect(lambda: self.custom.setVisible(self.preset.currentData() == "custom"))
        self.custom.setVisible(self.preset.currentData() == "custom")

        lbl = QLabel("Rarity")
        lbl.setObjectName("heading")
        lay.addWidget(lbl)
        names = book.meta.get("rarity")
        self.use_rarity = QCheckBox("Layers say how common they are in each square")
        self.use_rarity.setChecked(bool(names))
        lay.addWidget(self.use_rarity)
        self.rarity = QLineEdit(", ".join(names or ["Common", "Uncommon", "Rare"]))
        self.rarity.setToolTip("Up to three names, most common first")
        lay.addWidget(self.rarity)
        self.use_rarity.toggled.connect(self.rarity.setEnabled)
        self.rarity.setEnabled(bool(names))
        note = QLabel("Changing how time works keeps what it can; anything that doesn't fit becomes "
                      "\"all the time\". Turning rarity off keeps the levels (they come back if you turn it on).")
        note.setObjectName("faint")
        note.setWordWrap(True)
        lay.addWidget(note)
        lay.addStretch()
        box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        box.accepted.connect(self.accept)
        box.rejected.connect(self.reject)
        lay.addWidget(box)

    def result_meta(self):
        preset = self.preset.currentData()
        time = {"preset": preset}
        if preset == "custom":
            periods = [l.strip() for l in self.periods.toPlainText().splitlines() if l.strip()]
            groups = {}
            for line in self.groups.toPlainText().splitlines():
                if ":" in line:
                    g, members = line.split(":", 1)
                    ms = [m.strip() for m in members.split(",") if m.strip() in periods]
                    if g.strip() and ms:
                        groups[g.strip()] = ms
            time.update(periods=periods, groups=groups)
        rarity = None
        if self.use_rarity.isChecked():
            rarity = [n.strip() for n in self.rarity.text().split(",") if n.strip()][:3] or None
        return time, rarity
