"""The workbook behind Nous Atlas.

openpyxl's Workbook is the model: what you see is what gets saved, so the file
stays an ordinary .xlsx. This module adds what openpyxl doesn't do for us:
colour resolution (theme colours, tints), formula results, undo snapshots,
row/column insertion that keeps merges, sizes and formulas in step, and a
small hidden sheet where Atlas keeps its own settings (sheet groups, palette,
maps).
"""

import colorsys
import copy
import json
import os
import re
import zipfile
from xml.etree import ElementTree

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import atlas_formula as fx

META_SHEET = "_NousAtlas"
RANGES_SHEET = "Ranges"
DEFAULT_COL_WIDTH = 8.43          # Excel's default, in "characters"
DEFAULT_ROW_HEIGHT = 15.0         # points
PER_SHEET_KEYS = ("pinned_rows", "sheet_layout")   # meta entries keyed by sheet title


# ---------------------------------------------------------------- colours

_THEME_ORDER = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
                "accent4", "accent5", "accent6", "hlink", "folHlink"]

# Excel's legacy indexed colours (only the first 64 matter)
_INDEXED = [
    "000000", "FFFFFF", "FF0000", "00FF00", "0000FF", "FFFF00", "FF00FF", "00FFFF",
    "000000", "FFFFFF", "FF0000", "00FF00", "0000FF", "FFFF00", "FF00FF", "00FFFF",
    "800000", "008000", "000080", "808000", "800080", "008080", "C0C0C0", "808080",
    "9999FF", "993366", "FFFFCC", "CCFFFF", "660066", "FF8080", "0066CC", "CCCCFF",
    "000080", "FF00FF", "FFFF00", "00FFFF", "800080", "800000", "008080", "0000FF",
    "00CCFF", "CCFFFF", "CCFFCC", "FFFF99", "99CCFF", "FF99CC", "CC99FF", "FFCC99",
    "3366FF", "33CCCC", "99CC00", "FFCC00", "FF9900", "FF6600", "666699", "969696",
    "003366", "339966", "003300", "333300", "993300", "993366", "333399", "333333",
]


def theme_colors(wb):
    """The workbook theme's 12 colours as hex, in Excel's index order."""
    out = {}
    xml = getattr(wb, "loaded_theme", None)
    if xml:
        try:
            root = ElementTree.fromstring(xml)
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            scheme = root.find(".//a:clrScheme", ns)
            for child in scheme:
                tag = child.tag.split("}")[1]
                node = child[0]
                val = node.get("lastClr") or node.get("val")
                if val:
                    out[tag] = val.upper()
        except Exception:
            out = {}
    fallback = {"lt1": "FFFFFF", "dk1": "000000", "lt2": "E7E6E6", "dk2": "44546A",
                "accent1": "4472C4", "accent2": "ED7D31", "accent3": "A5A5A5",
                "accent4": "FFC000", "accent5": "5B9BD5", "accent6": "70AD47",
                "hlink": "0563C1", "folHlink": "954F72"}
    return [out.get(k, fallback[k]) for k in _THEME_ORDER]


def apply_tint(hex6, tint):
    if not tint:
        return hex6
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))


def resolve_color(color, theme):
    """openpyxl Color -> 'RRGGBB' or None (automatic / unset)."""
    if color is None:
        return None
    try:
        kind = color.type
        if kind == "rgb":
            rgb = color.rgb
            if not isinstance(rgb, str) or len(rgb) < 6:
                return None
            return rgb[-6:].upper()
        if kind == "theme":
            idx = color.theme
            if idx is None or idx >= len(theme):
                return None
            return apply_tint(theme[idx], color.tint or 0)
        if kind == "indexed":
            idx = color.indexed
            if idx is not None and idx < len(_INDEXED):
                return apply_tint(_INDEXED[idx], color.tint or 0)
            return None                     # 64/65 = system foreground/background
    except Exception:
        return None
    return None


# ---------------------------------------------------------------- values


_INT_RE = re.compile(r"[-+]?\d+")
_FLOAT_RE = re.compile(r"[-+]?(\d+\.\d*|\.\d+|\d+)([eE][-+]?\d+)?")
_PCT_RE = re.compile(r"[-+]?(\d+\.?\d*|\.\d+)%")


def parse_input(text):
    """What a typed string becomes in the cell: (value, number_format or None)."""
    if text == "":
        return None, None
    if text.startswith("'"):
        return text[1:], "@"
    if text.startswith("=") and len(text) > 1:
        return text, None
    s = text.strip()
    if _INT_RE.fullmatch(s) and len(s.lstrip("+-")) < 16:
        return int(s), None
    if _FLOAT_RE.fullmatch(s):
        return float(s), None
    if _PCT_RE.fullmatch(s):
        return float(s[:-1]) / 100, "0%"
    if s.upper() in ("TRUE", "FALSE"):
        return s.upper() == "TRUE", None
    return text, None


_NUMFMT_DECIMALS = re.compile(r"^(#,##)?0(\.(0+))?(%?)$")


def format_value(value, number_format):
    """Display text for a computed value."""
    if value is None:
        return ""
    if isinstance(value, fx.XlError):
        return str(value)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        m = _NUMFMT_DECIMALS.match(number_format or "")
        if m:
            decimals = len(m.group(3) or "")
            v = float(value) * (100 if m.group(4) else 1)
            v = fx.excel_round(v, decimals)
            text = f"{v:,.{decimals}f}" if m.group(1) else f"{v:.{decimals}f}"
            return text + m.group(4)
        return fx.format_general(value)
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ") if hasattr(value, "hour") else value.isoformat()
    return str(value)


# ---------------------------------------------------------------- snapshots


def _copy_cell(cell):
    new = copy.copy(cell)
    if cell._style is not None:
        new._style = copy.copy(cell._style)
    return new


class SheetState:
    """Everything about one sheet that an edit can change."""

    def __init__(self, ws):
        self.ws = ws
        self.title = ws.title
        self.cells = {k: _copy_cell(c) for k, c in ws._cells.items()}
        self.merged = [str(r) for r in ws.merged_cells.ranges]
        self.rows = {k: (d.height, d.hidden, d.customHeight) for k, d in ws.row_dimensions.items()}
        self.cols = {k: (d.width, d.hidden, d.customWidth) for k, d in ws.column_dimensions.items()}
        self.freeze = ws.freeze_panes
        self.state = ws.sheet_state

    def restore(self):
        ws = self.ws
        ws.title = self.title
        ws._cells = {k: _copy_cell(c) for k, c in self.cells.items()}
        for c in ws._cells.values():
            c.parent = ws
        ws.merged_cells.ranges = set()
        for rng in self.merged:
            ws.merged_cells.add(rng)
        ws.row_dimensions.clear()
        for k, (h, hidden, custom) in self.rows.items():
            d = ws.row_dimensions[k]
            d.height, d.hidden = h, hidden
        ws.column_dimensions.clear()
        for k, (w, hidden, custom) in self.cols.items():
            d = ws.column_dimensions[k]
            if w is not None:
                d.width = w
            d.hidden = hidden
        ws.freeze_panes = self.freeze
        ws.sheet_state = self.state
        rebuild_hyperlinks(ws)


class Snapshot:
    def __init__(self, book, sheets):
        self.order = list(book.wb._sheets)
        self.active = book.wb.active.title if book.wb.active else None
        self.sheets = [SheetState(ws) for ws in sheets]
        self.meta = copy.deepcopy(book.meta)
        self.ranges = copy.deepcopy(book.ranges) if book.ranges is not None else None
        self.named = [(ns.name, copy.copy(ns.font), copy.copy(ns.fill), copy.copy(ns.border),
                       copy.copy(ns.alignment), ns.number_format) for ns in book.wb._named_styles]

    def restore(self, book):
        book.wb._sheets = list(self.order)
        for s in self.sheets:
            s.restore()
        book.meta = copy.deepcopy(self.meta)
        book.ranges = copy.deepcopy(self.ranges) if self.ranges is not None else None
        for name, font, fill, border, align, numfmt in self.named:
            for ns in book.wb._named_styles:
                if ns.name == name:
                    ns.font, ns.fill, ns.border, ns.alignment = font, fill, border, align
                    ns.number_format = numfmt


def rebuild_hyperlinks(ws):
    links = []
    for cell in ws._cells.values():
        link = getattr(cell, "_hyperlink", None)
        if link is not None:
            link.ref = cell.coordinate
            links.append(link)
    ws._hyperlinks = links


# ---------------------------------------------------------------- the book


class Book:
    """A workbook open in Atlas."""

    def __init__(self, wb=None, path=None):
        self.wb = wb or openpyxl.Workbook()
        self.path = path
        self.warnings = []          # what the import couldn't carry over
        self.must_save_as = False   # saving over the original would lose things
        self.theme = theme_colors(self.wb)
        self.meta = {}
        self.ranges = None          # filled in by atlas_ranges
        self.undo_stack, self.redo_stack = [], []
        self.dirty = False
        self.calc = fx.Calc(self.raw, lambda: self.wb.sheetnames)
        self.listeners = []         # called after any change: fn(kind)
        for ws in self.wb.worksheets:
            normalise_columns(ws)
        self._load_meta()

    # -- opening and saving
    @classmethod
    def open(cls, path):
        wb = openpyxl.load_workbook(path, rich_text=False)
        book = cls(wb, path)
        book.warnings, book.must_save_as = inspect_package(path, wb)
        return book

    def save(self, path=None):
        path = path or self.path
        self._write_meta()
        tmp = path + ".nousatlas-tmp"
        self.wb.save(tmp)
        try:
            os.replace(tmp, path)
        except OSError:
            try:
                os.remove(tmp)
            except OSError:
                pass
            raise
        self.path = path
        self.must_save_as = False
        self.dirty = False
        self.notify("saved")

    def _load_meta(self):
        if META_SHEET in self.wb.sheetnames:
            ws = self.wb[META_SHEET]
            text = "".join(str(ws.cell(r, 1).value or "") for r in range(1, ws.max_row + 1))
            try:
                self.meta = json.loads(text) if text else {}
            except ValueError:
                self.meta = {}
        self.meta.setdefault("groups", [])
        self.meta.setdefault("maps", {})
        self.meta.setdefault("palette", None)
        self.meta.setdefault("layer_colors", {})

    def _write_meta(self):
        if self.ranges_writer:
            self.ranges_writer(self)
        text = json.dumps(self.meta, ensure_ascii=False)
        if META_SHEET in self.wb.sheetnames:
            ws = self.wb[META_SHEET]
            ws.delete_rows(1, ws.max_row)
        else:
            ws = self.wb.create_sheet(META_SHEET)
        ws.sheet_state = "veryHidden"
        # Excel caps a cell at 32,767 characters
        for i in range(0, max(1, len(text)), 32000):
            ws.cell(i // 32000 + 1, 1).value = text[i:i + 32000]
        if self.wb.active is ws or self.wb.active is None or self.wb.active.sheet_state != "visible":
            for i, s in enumerate(self.wb.worksheets):
                if s.sheet_state == "visible":
                    self.wb.active = i
                    break

    ranges_writer = None            # set by atlas_ranges: writes the Ranges sheet

    # -- sheets the user sees
    def user_sheets(self):
        return [ws for ws in self.wb.worksheets
                if ws.title != META_SHEET and ws.sheet_state != "veryHidden"]

    def sheet(self, title):
        return self.wb[title] if title in self.wb.sheetnames else None

    # -- values
    def raw(self, sheet, row, col):
        ws = self.wb[sheet] if sheet in self.wb.sheetnames else None
        if ws is None:
            return None
        cell = ws._cells.get((row, col))
        return None if cell is None else cell.value

    def value(self, ws, row, col):
        return self.calc.value(ws.title, row, col)

    def display_text(self, ws, row, col):
        cell = ws._cells.get((row, col))
        if cell is None or cell.value is None:
            return ""
        v = self.calc.value(ws.title, row, col) if fx.is_formula(cell.value) else cell.value
        return format_value(v, cell.number_format)

    @staticmethod
    def edit_text(ws, row, col):
        cell = ws._cells.get((row, col))
        if cell is None or cell.value is None:
            return ""
        v = cell.value
        if isinstance(v, str):
            return ("'" + v) if cell.number_format == "@" and _looks_numeric(v) else v
        return format_value(v, "General") if isinstance(v, (int, float, bool)) else str(v)

    # -- change tracking and undo
    def notify(self, kind):
        for fn in list(self.listeners):
            fn(kind)

    def begin(self, sheets=None):
        """Call before changing anything. `sheets`: worksheets that will change
        (None = all of them)."""
        sheets = self.wb.worksheets if sheets is None else sheets
        self.undo_stack.append(Snapshot(self, sheets))
        del self.undo_stack[:-200]
        self.redo_stack.clear()

    def done(self, kind="edit"):
        self.dirty = True
        self.calc.invalidate()
        self.notify(kind)

    def undo(self):
        return self._swap(self.undo_stack, self.redo_stack)

    def redo(self):
        return self._swap(self.redo_stack, self.undo_stack)

    def _swap(self, src, dst):
        if not src:
            return False
        snap = src.pop()
        dst.append(Snapshot(self, [s.ws for s in snap.sheets]))
        snap.restore(self)
        self.dirty = True
        self.calc.invalidate()
        self.notify("undo")
        return True

    # -- cells
    def set_value(self, ws, row, col, text):
        """Store typed text (already inside begin/done)."""
        value, numfmt = parse_input(text)
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            return
        cell.value = value
        if numfmt:
            cell.number_format = numfmt
        elif cell.number_format == "@" and not text.startswith("'"):
            cell.number_format = "General"

    def cells_in(self, ws, r1, c1, r2, c2, create=True):
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                cell = ws.cell(r, c) if create else ws._cells.get((r, c))
                if cell is not None:
                    yield cell

    # -- structure
    def insert(self, ws, axis, at, n):
        self._restructure(ws, axis, at, n)

    def delete(self, ws, axis, at, n):
        self._restructure(ws, axis, at, -n)

    def _restructure(self, ws, axis, at, delta):
        """Insert (delta > 0) or delete (delta < 0) rows/columns starting at
        `at`, keeping merges, sizes, frozen panes, hyperlinks and every
        formula in the workbook in step."""
        merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
        for m in list(ws.merged_cells.ranges):
            ws.merged_cells.remove(m)
        for key in [k for k, c in ws._cells.items() if isinstance(c, MergedCell)]:
            del ws._cells[key]

        if axis == "row":
            if delta > 0:
                ws.insert_rows(at, delta)
            else:
                ws.delete_rows(at, -delta)
        else:
            if delta > 0:
                ws.insert_cols(at, delta)
            else:
                ws.delete_cols(at, -delta)

        def shift(lo, hi):
            if delta > 0:
                return (lo + delta if lo >= at else lo, hi + delta if hi >= at else hi)
            gone_hi = at - delta - 1
            if lo >= at and hi <= gone_hi:
                return None
            nlo = at if at <= lo <= gone_hi else (lo + delta if lo > gone_hi else lo)
            nhi = at - 1 if at <= hi <= gone_hi else (hi + delta if hi > gone_hi else hi)
            return (nlo, nhi)

        for r1, c1, r2, c2 in merges:
            span = shift(r1, r2) if axis == "row" else shift(c1, c2)
            if span is None:
                continue
            if axis == "row":
                r1, r2 = span
            else:
                c1, c2 = span
            if (r1, c1) != (r2, c2):
                ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

        # row heights / column widths move with their rows / columns
        if axis == "row":
            old = {k: (d.height, d.hidden, d.customHeight) for k, d in ws.row_dimensions.items()}
            ws.row_dimensions.clear()
            for k, v in old.items():
                span = shift(k, k)
                if span is None:
                    continue
                d = ws.row_dimensions[span[0]]
                d.height, d.hidden = v[0], v[1]
        else:
            old = {}
            for k, d in ws.column_dimensions.items():
                old[d.min or openpyxl.utils.column_index_from_string(k)] = (d.width, d.hidden, d.customWidth)
            ws.column_dimensions.clear()
            for idx, v in old.items():
                span = shift(idx, idx)
                if span is None:
                    continue
                d = ws.column_dimensions[get_column_letter(span[0])]
                d.width = v[0]
                d.hidden = v[1]

        # frozen panes
        if ws.freeze_panes:
            m = re.fullmatch(r"([A-Z]+)(\d+)", ws.freeze_panes)
            if m:
                fc = openpyxl.utils.column_index_from_string(m.group(1))
                fr = int(m.group(2))
                if axis == "row" and at < fr:
                    fr = max(1, fr + delta if delta > 0 else fr - min(-delta, fr - at))
                if axis == "col" and at < fc:
                    fc = max(1, fc + delta if delta > 0 else fc - min(-delta, fc - at))
                ws.freeze_panes = None if (fr, fc) == (1, 1) else f"{get_column_letter(fc)}{fr}"

        # formulas everywhere
        for other in self.wb.worksheets:
            for cell in other._cells.values():
                if fx.is_formula(cell.value):
                    cell.value = fx.shift_refs(cell.value, other.title, ws.title, axis, at, delta)

        # rows sized by hand move with their rows
        if axis == "row" and ws.title in self.meta.get("pinned_rows", {}):
            moved = []
            for r in self.meta["pinned_rows"][ws.title]:
                span = shift(r, r)
                if span:
                    moved.append(span[0])
            self.meta["pinned_rows"][ws.title] = sorted(set(moved))

        # maps whose grid sits after the change move with it
        info = self.meta["maps"].get(ws.title)
        if info:
            key = "row" if axis == "row" else "col"
            span = shift(info[key], info[key])
            if span:
                info[key] = span[0]
        rebuild_hyperlinks(ws)

    def merge(self, ws, r1, c1, r2, c2):
        for m in list(ws.merged_cells.ranges):
            if not (m.max_row < r1 or m.min_row > r2 or m.max_col < c1 or m.min_col > c2):
                ws.unmerge_cells(str(m))
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    def unmerge(self, ws, r1, c1, r2, c2):
        found = False
        for m in list(ws.merged_cells.ranges):
            if not (m.max_row < r1 or m.min_row > r2 or m.max_col < c1 or m.min_col > c2):
                ws.unmerge_cells(str(m))
                found = True
        return found

    def merged_at(self, ws, row, col):
        for m in ws.merged_cells.ranges:
            if m.min_row <= row <= m.max_row and m.min_col <= col <= m.max_col:
                return m
        return None

    # -- sheets
    def unique_title(self, base):
        title, n = base[:31], 2
        while title.lower() in (s.lower() for s in self.wb.sheetnames):
            suffix = f" ({n})"
            title, n = base[:31 - len(suffix)] + suffix, n + 1
        return title

    def add_sheet(self, title, index=None):
        ws = self.wb.create_sheet(self.unique_title(title), index)
        return ws

    def rename_sheet(self, ws, new):
        old = ws.title
        ws.title = new
        for other in self.wb.worksheets:
            for cell in other._cells.values():
                if fx.is_formula(cell.value):
                    cell.value = fx.rename_sheet_refs(cell.value, old, new)
        for group in self.meta["groups"]:
            group["sheets"] = [new if s == old else s for s in group["sheets"]]
        if old in self.meta["maps"]:
            self.meta["maps"][new] = self.meta["maps"].pop(old)
        for key in PER_SHEET_KEYS:
            if old in self.meta.get(key, {}):
                self.meta[key][new] = self.meta[key].pop(old)
        if self.ranges is not None:
            self.ranges.rename_map(old, new)

    def delete_sheet(self, ws):
        title = ws.title
        self.wb.remove(ws)
        for other in self.wb.worksheets:
            for cell in other._cells.values():
                if fx.is_formula(cell.value):
                    cell.value = fx._rewrite_refs(
                        cell.value, lambda ref: fx.REF if ref.sheet and ref.sheet.lower() == title.lower() else None)
        for group in self.meta["groups"]:
            group["sheets"] = [s for s in group["sheets"] if s != title]
        self.meta["maps"].pop(title, None)
        for key in PER_SHEET_KEYS:
            self.meta.get(key, {}).pop(title, None)
        if self.ranges is not None:
            self.ranges.drop_map(title)

    def duplicate_sheet(self, ws):
        new = self.wb.copy_worksheet(ws)
        new.title = self.unique_title(ws.title + " copy")
        new.freeze_panes = ws.freeze_panes
        self.wb.move_sheet(new, self.wb.worksheets.index(ws) + 1 - self.wb.worksheets.index(new))
        for group in self.meta["groups"]:
            if ws.title in group["sheets"]:
                group["sheets"].insert(group["sheets"].index(ws.title) + 1, new.title)
        if ws.title in self.meta["maps"]:
            self.meta["maps"][new.title] = dict(self.meta["maps"][ws.title])
        for key in PER_SHEET_KEYS:
            if ws.title in self.meta.get(key, {}):
                self.meta[key][new.title] = copy.deepcopy(self.meta[key][ws.title])
        return new

    # -- colours used anywhere (for the palette tools)
    def colors_in_use(self):
        """{'RRGGBB': count} over font colours and fills."""
        counts = {}
        for ws in self.user_sheets():
            for cell in ws._cells.values():
                if cell.has_style:
                    for col in (cell.font.color if cell.font else None,
                                cell.fill.fgColor if cell.fill and cell.fill.fill_type else None):
                        hx = resolve_color(col, self.theme)
                        if hx:
                            counts[hx] = counts.get(hx, 0) + 1
        return counts

    def replace_colors(self, mapping):
        """Swap colours everywhere: mapping {'RRGGBB': 'RRGGBB'}."""
        mapping = {k.upper(): v.upper() for k, v in mapping.items() if k.upper() != v.upper()}
        if not mapping:
            return 0
        n = 0
        for ws in self.user_sheets():
            for cell in ws._cells.values():
                if not cell.has_style:
                    continue
                fc = resolve_color(cell.font.color, self.theme) if cell.font else None
                if fc in mapping:
                    f = copy.copy(cell.font)
                    f.color = "FF" + mapping[fc]
                    cell.font = f
                    n += 1
                if cell.fill and cell.fill.fill_type:
                    bc = resolve_color(cell.fill.fgColor, self.theme)
                    if bc in mapping:
                        cell.fill = PatternFill("solid", fgColor="FF" + mapping[bc])
                        n += 1
                b = cell.border
                if b is not None and any(s is not None and s.style for s in (b.left, b.right, b.top, b.bottom)):
                    sides = {}
                    changed = False
                    for name in ("left", "right", "top", "bottom"):
                        s = getattr(b, name)
                        sc = resolve_color(s.color, self.theme) if s is not None and s.style else None
                        if sc in mapping:
                            sides[name] = Side(style=s.style, color="FF" + mapping[sc])
                            changed = True
                        else:
                            sides[name] = copy.copy(s) if s is not None else Side()
                    if changed:
                        cell.border = Border(**sides)
        for ns in self.wb._named_styles:
            fc = resolve_color(ns.font.color, self.theme) if ns.font else None
            if fc in mapping:
                f = copy.copy(ns.font)
                f.color = "FF" + mapping[fc]
                ns.font = f
            if ns.fill and getattr(ns.fill, "fill_type", None):
                bc = resolve_color(ns.fill.fgColor, self.theme)
                if bc in mapping:
                    ns.fill = PatternFill("solid", fgColor="FF" + mapping[bc])
        return n


def _looks_numeric(s):
    s = s.strip()
    return bool(_INT_RE.fullmatch(s) or _FLOAT_RE.fullmatch(s) or _PCT_RE.fullmatch(s))


def normalise_columns(ws):
    """openpyxl keeps <col min=2 max=18 width=6> as one entry keyed 'B'.
    Split those spans so every column has its own entry."""
    spans = []
    for key, d in list(ws.column_dimensions.items()):
        lo = d.min or openpyxl.utils.column_index_from_string(key)
        hi = d.max or lo
        if hi > lo and hi - lo < 16384:
            spans.append((lo, min(hi, 1000), d.width, d.hidden, d.customWidth))
    for lo, hi, width, hidden, custom in spans:
        for idx in range(lo, hi + 1):
            d = ws.column_dimensions[get_column_letter(idx)]
            d.min = d.max = idx
            d.width = width
            d.hidden = hidden


def inspect_package(path, wb):
    """What's in the file that Atlas can't keep or show. Returns
    (warnings, must_save_as)."""
    warnings, lossy = [], False
    try:
        names = zipfile.ZipFile(path).namelist()
    except (zipfile.BadZipFile, OSError):
        return warnings, lossy
    def has(prefix):
        return any(n.startswith(prefix) for n in names)
    if has("xl/charts/"):
        warnings.append("Charts aren't kept. Saving over this file would remove them.")
        lossy = True
    if has("xl/media/") or has("xl/drawings/drawing"):
        warnings.append("Pictures and drawings aren't kept. Saving over this file would remove them.")
        lossy = True
    if any("vbaProject" in n for n in names):
        warnings.append("Macros aren't kept.")
        lossy = True
    if has("xl/threadedComments/"):
        warnings.append("Threaded comments come in as plain notes; replies are flattened.")
    if has("xl/slicers/") or has("xl/timelines/"):
        warnings.append("Slicers and timelines aren't kept.")
        lossy = True
    for ws in wb.worksheets:
        if ws.conditional_formatting and len(ws.conditional_formatting):
            warnings.append(f"{ws.title}: conditional formatting is kept in the file but not shown.")
        if ws.data_validations.dataValidation:
            warnings.append(f"{ws.title}: dropdown lists and input rules are kept but not enforced.")
    return warnings, lossy


# ---------------------------------------------------------------- styling helpers


def set_font(cell, **changes):
    f = copy.copy(cell.font) if cell.font else Font()
    for k, v in changes.items():
        setattr(f, k, v)
    cell.font = f


def set_alignment(cell, **changes):
    a = copy.copy(cell.alignment) if cell.alignment else Alignment()
    for k, v in changes.items():
        setattr(a, k, v)
    cell.alignment = a


def set_fill(cell, hex6):
    cell.fill = PatternFill(fill_type=None) if hex6 is None else PatternFill("solid", fgColor="FF" + hex6)


def set_borders(cell, left=None, right=None, top=None, bottom=None, keep=True):
    """Each side: True = thin line, False = none, None = leave as is."""
    b = cell.border or Border()
    sides = {}
    for name, val in (("left", left), ("right", right), ("top", top), ("bottom", bottom)):
        cur = getattr(b, name)
        if val is None:
            sides[name] = copy.copy(cur) if cur is not None else Side()
        elif val is False:
            sides[name] = Side()
        else:
            sides[name] = Side(style="thin", color=val if isinstance(val, str) else None)
    cell.border = Border(**sides)
