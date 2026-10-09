"""Makes sample/Gullwing-Isle.xlsx, the made-up sample world that ships with
the app (Help > Open the sample world) and is linked from the download page.
Run from the project folder, then commit the .xlsx:
    .venv\\Scripts\\python.exe make_sample.py"""

import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.hyperlink import Hyperlink

import atlas_ranges as R
from atlas_model import Book

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "sample", "Gullwing-Isle.xlsx")
MAP = "Island Map"

# ------------------------------------------------------------------ the island
TERRAIN = {   # code: (name, fill)
    "~": ("Sea", "5E81AC"),
    ".": ("Beach", "EBCB8B"),
    "c": ("Cliffs", "7B8394"),
    "h": ("Hills", "B9A07E"),
    "m": ("Meadow", "A3BE8C"),
    "f": ("Forest", "5B8A5E"),
    "w": ("Marsh", "8FBCBB"),
    "v": ("Village", "D08770"),
}
GRID = [
    "~~~~~~~~~~~~~~",
    "~~~..cchhc~~~~",
    "~~.mmfhhhc.~~~",
    "~.mmmffhhff.~~",
    "~.mvvmfffff.~~",
    "~~.vvmmfffw.~~",
    "~~..mmmmwww.~~",
    "~~~...mmww..~~",
    "~~~~~~...~~~~~",
    "~~~~~~~~~~~~~~",
]
TOP, LEFT = 3, 2          # the letters row and the numbers column

HEAD_FONT = Font(name="Segoe UI", size=11, bold=True, color="ECEFF4")
HEAD_FILL = PatternFill("solid", fgColor="3B4252")
TITLE_FONT = Font(name="Segoe UI", size=18, bold=True)
BODY = Font(name="Segoe UI", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")


def sq(label):
    """'E8' -> (5, 8)"""
    return (ord(label[0]) - 64, int(label[1:]))


def map_ref(label):
    x, y = sq(label)
    col = openpyxl.utils.get_column_letter(LEFT + x)
    return f"'{MAP}'!{col}{TOP + y}"


def table(ws, headers, rows, widths):
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24
    for r, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            c.font, c.alignment = BODY, WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def link(cell, location):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=location)


wb = openpyxl.Workbook()

# ------------------------------------------------------------------ Start Here
ws = wb.active
ws.title = "Start Here"
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 92
ws["B2"] = "Welcome to Gullwing Isle"
ws["B2"].font = TITLE_FONT
ws.row_dimensions[2].height = 34
ws.row_dimensions[3].height = 40
ws.row_dimensions[5].height = 28
ws["B3"] = ("A tiny made-up world to try Nous Atlas on. This is your own copy, so change anything you like. "
            "Ctrl+Z undoes.")
ws["B3"].font, ws["B3"].alignment = Font(name="Segoe UI", size=11, italic=True, color="7B8394"), WRAP
steps = [
    ("Things to try", None),
    ("1.  The sheets are listed on the left, in groups. Click one to open it.", None),
    ("2.  Open the Island Map, and rest the mouse on a square. A card shows what lives and happens there.", MAP + "!A1"),
    ("3.  Click Layers on the toolbar. Each creature is a layer; click an eye to hide or show it.", None),
    ("4.  Squares split in two show day and night: the Red fox is common at night but rare by day.", None),
    ("5.  On the Creatures sheet, right-click “Sand crabs” and choose Make a map layer. Then pick up the "
     "brush and drag over the beaches.", "Creatures!A1"),
    ("6.  Places are linked to the map. On the Places sheet, click a square's button to jump to it on the map.",
     "Places!A1"),
    ("7.  Press Ctrl+F and type “moth” to search every sheet at once.", None),
    ("8.  Press Ctrl+P and type “peo” to jump to the People sheet. Alt+← goes back.", None),
    ("9.  Make it yours: type over anything, add a sheet (Insert → Sheet), or start a new world with "
     "File → New workbook.", None),
]
r = 5
for text, target in steps:
    c = ws.cell(r, 2, text)
    if target is None and r == 5:
        c.font = Font(name="Segoe UI", size=13, bold=True)
    else:
        c.font, c.alignment = BODY, WRAP
    if r > 5:
        ws.row_dimensions[r].height = 24 if len(text) < 100 else 42
    if target:
        sheet, _, ref = target.partition("!")
        link(c, f"'{sheet}'!{ref}")
    r += 1
ws.cell(r + 1, 2, "Your file stays an ordinary Excel workbook (.xlsx). Excel and Google Sheets can open it too; "
                  "they just won't show the map layers.").font = Font(name="Segoe UI", size=10, color="7B8394")
ws.cell(r + 1, 2).alignment = WRAP

# ------------------------------------------------------------------ the map
m = wb.create_sheet(MAP)
m.column_dimensions["A"].width = 3
m.column_dimensions[openpyxl.utils.get_column_letter(LEFT)].width = 4
for x in range(1, 15):
    m.column_dimensions[openpyxl.utils.get_column_letter(LEFT + x)].width = 5.2
m.row_dimensions[TOP].height = 20
m.cell(1, LEFT, "Gullwing Isle").font = Font(name="Segoe UI", size=16, bold=True)
m.row_dimensions[1].height = 30
thin = Side(style="thin", color="4C566A")
for x in range(1, 15):
    c = m.cell(TOP, LEFT + x, chr(64 + x))
    c.font, c.alignment = Font(name="Segoe UI", size=10, bold=True), Alignment(horizontal="center")
for y, row in enumerate(GRID, 1):
    m.row_dimensions[TOP + y].height = 30
    c = m.cell(TOP + y, LEFT, y)
    c.font, c.alignment = Font(name="Segoe UI", size=10, bold=True), Alignment(horizontal="center", vertical="center")
    for x, code in enumerate(row, 1):
        c = m.cell(TOP + y, LEFT + x)
        c.fill = PatternFill("solid", fgColor=TERRAIN[code][1])
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
# a few names on the map itself
for label, text, size in (("E8", "⌂", 16), ("D5", "Mossy", 9)):
    x, y = sq(label)
    c = m.cell(TOP + y, LEFT + x, text)
    c.font = Font(name="Segoe UI", size=size, bold=True, color="2E3440")
    c.alignment = Alignment(horizontal="center", vertical="center")
# the key
kc = LEFT + 16
m.column_dimensions[openpyxl.utils.get_column_letter(kc)].width = 4
m.column_dimensions[openpyxl.utils.get_column_letter(kc + 1)].width = 14
m.cell(TOP, kc, "KEY").font = Font(name="Segoe UI", size=10, bold=True)
for i, (name, fill) in enumerate(TERRAIN.values(), 1):
    m.cell(TOP + i, kc).fill = PatternFill("solid", fgColor=fill)
    m.cell(TOP + i, kc).border = Border(left=thin, right=thin, top=thin, bottom=thin)
    m.cell(TOP + i, kc + 1, name).font = BODY

# ------------------------------------------------------------------ Places
places = [
    ("Gullwing Lighthouse", "E8", "A white tower on the south beach. The lamp is lit every night at dusk."),
    ("Mossy", "D5", "The only village: twelve cottages, a bakery, a harbour wall and a lot of cats."),
    ("Puffin Point", "F2", "Sea cliffs where the puffins nest from spring to late summer."),
    ("Whisper Wood", "H5", "Old oaks and ferns. People say the trees talk when the wind is from the east."),
    ("Fen of Lights", "J7", "A reedy marsh. On summer nights the glow moths rise from it in clouds."),
    ("Silver Brook", "H3", "Starts as a spring in the hills and runs south to the sea."),
    ("Brookmouth", "H9", "Where the brook meets the sea. Good for paddling and finding sea glass."),
    ("The Old Wreck", "L9", "A fishing boat, wrecked long ago. At low tide you can almost walk to it."),
]
p = wb.create_sheet("Places")
table(p, ["Place", "Square", "What's there"], [(n, s, d) for n, s, d in places], [24, 10, 70])
for i, (_, label, _) in enumerate(places, 2):
    link(p.cell(i, 2), map_ref(label))

# ------------------------------------------------------------------ Creatures
c = wb.create_sheet("Creatures")
table(c, ["Creature", "Where", "When", "Notes"], [
    ("Puffins", "Puffin Point and the sea around it", "Day", "Nest in burrows at the cliff top. Fish all day, sleep on the water."),
    ("Glow moths", "Fen of Lights, the edge of Whisper Wood", "Night",
     "Pale green light. Juniper says they only glow when it's warm."),
    ("Red fox", "Meadows and woods, all over", "Mostly night",
     "Shy by day. Steals from the bakery bins at night."),
    ("Sand crabs", "Every beach", "Day and night", "Not on the map yet: try making a layer for them."),
    ("Grey seal", "Off the Old Wreck", "Day", "One old seal the children call Biscuit."),
], [16, 34, 14, 60])

# ------------------------------------------------------------------ People
pe = wb.create_sheet("People")
table(pe, ["Name", "Lives at", "Job", "Notes"], [
    ("Mira Holt", "Gullwing Lighthouse", "Lighthouse keeper", "Keeps the logbook of every ship and storm. Never misses a dusk."),
    ("Old Tam", "Mossy", "Ferryman", "Rows to the mainland on Tuesdays, weather allowing. Tells tall tales for free."),
    ("Pip and Wren Pell", "Mossy", "Bakers", "Twins. Pip bakes, Wren talks. Their seed cake is famous."),
    ("Juniper", "Edge of Whisper Wood", "Healer", "Gathers moss and marsh herbs. Knows the moths better than anyone."),
], [18, 22, 18, 60])
for i, place in ((2, "Gullwing Lighthouse"), (3, "Mossy"), (4, "Mossy")):
    row = next(n for n, (pl, _, _) in enumerate(places, 2) if pl == place)
    link(pe.cell(i, 2), f"'Places'!A{row}")

# ------------------------------------------------------------------ Ideas
idea = wb.create_sheet("Ideas")
table(idea, ["Idea", "Notes"], [
    ("The lamp goes out", "One night the lighthouse lamp won't light. Mira swears she filled it."),
    ("Moth year", "Every seventh summer the moths come out in their thousands."),
    ("A visitor", "A stranger arrives on Tam's ferry, asking about the wreck."),
    ("Your idea here", "Type over this row, or add more below."),
], [24, 76])

# ------------------------------------------------------------------ into Atlas: maps, groups, layers, notes
book = Book(wb)
book.meta["time"] = {"preset": "daynight"}
book.meta["rarity"] = ["Common", "Uncommon", "Rare"]
book.meta["maps"][MAP] = R.detect_map(m)
assert book.meta["maps"][MAP] == {"row": TOP + 1, "col": LEFT + 1, "rows": 10, "cols": 14}, book.meta["maps"][MAP]
book.meta["groups"] = [
    {"name": "Welcome", "sheets": ["Start Here"], "collapsed": False},
    {"name": "The island", "sheets": [MAP, "Places"], "collapsed": False},
    {"name": "Notes", "sheets": ["Creatures", "People", "Ideas"], "collapsed": False},
]
book.meta["scanned"] = True
book.meta["last_map"] = MAP
book.meta["square_notes"] = {MAP: {
    "E8": "Mira's logbook says the lamp has been lit 9,412 nights in a row.",
    "L9": "At low tide you can almost walk out to the wreck. Don't.",
    "J7": "Best on a warm, moonless night.",
    "H3": "The spring. The water is cold enough to make your teeth ache.",
}}
scheme = R.TimeScheme.from_meta(book.meta)
book.ranges = R.Ranges(scheme)
DAY, NIGHT = {"day"}, {"night"}
BOTH = DAY | NIGHT


def where(codes):
    return [(x, y) for y, row in enumerate(GRID, 1) for x, ch in enumerate(row, 1) if ch in codes]


def near(squares, codes):
    """Squares of the given terrain next to any of `squares`."""
    s = set(squares)
    return [q for q in where(codes) if q not in s and any(abs(q[0] - a) <= 1 and abs(q[1] - b) <= 1 for a, b in s)]


rs = book.ranges
puffins = rs.add("Puffins", "ECEFF4")
rs.paint(puffins, MAP, DAY, where("c"), 3)
rs.paint(puffins, MAP, DAY, near(where("c"), "~"), 2)
moths = rs.add("Glow moths", "A3E635")
rs.paint(moths, MAP, NIGHT, where("w"), 3)
rs.paint(moths, MAP, NIGHT, near(where("w"), "f"), 2)
fox = rs.add("Red fox", "D9622B")
rs.paint(fox, MAP, NIGHT, [sq(s) for s in ("E3", "F4", "F5", "E7", "F7", "G6", "G7")], 3)
rs.paint(fox, MAP, DAY, [sq(s) for s in ("F4", "G6")], 1)
seal = rs.add("Grey seal", "9AA5B8")
seal.style = "marker"
rs.paint(seal, MAP, BOTH, [sq("L9")], 3)
brook = rs.add("Silver Brook", "88C0D0")
brook.style = "path"
rs.paint(brook, MAP, BOTH, [sq(s) for s in ("H3", "H4", "H5", "H6", "H7", "H8", "H9")], 3)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
book.save(OUT)
print("Made", OUT)
