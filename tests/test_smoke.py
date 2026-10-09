r"""Opens a small made-up workbook in the real (offscreen) window, edits a cell,
saves and reopens it. Needs no private files, so GitHub's build machines run it.

Run from the project folder:  .venv\Scripts\python.exe tests\test_smoke.py"""
import os, sys, tempfile
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["NOUSATLAS_SETTINGS"] = "NousAtlasSmoke"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import faulthandler
faulthandler.dump_traceback_later(120, exit=True)      # never hang on a stray dialog
from PySide6.QtWidgets import QApplication, QMessageBox
import openpyxl
import nousatlas

QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.Yes)
QMessageBox.information = staticmethod(lambda *a, **k: None)
QMessageBox.warning = staticmethod(lambda *a, **k: None)

here = tempfile.mkdtemp(prefix="nousatlas-smoke-")
path = os.path.join(here, "smoke.xlsx")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Places"
ws.append(["Name", "Notes"])
ws.append(["Old Well", "Dry since spring"])
ws.append(["Red Arch", "=LEN(B2)"])
ws.merge_cells("A5:B5")
ws["A5"] = "A merged heading"
ws["A2"].hyperlink = "#'Fauna'!A1"
f = wb.create_sheet("Fauna")
f.append(["Creature", "Where"])
f.append(["Sand hare", "Old Well"])
wb.save(path)

app = QApplication(sys.argv)
w = nousatlas.MainWindow(); w.resize(1200, 800); w.show()
assert w.open_path(path)
assert [s.title for s in w.sidebar_order()] == ["Places", "Fauna"], [s.title for s in w.sidebar_order()]
w.open_sheet("Fauna")
app.processEvents()
book = w.book
book.begin([book.sheet("Places")])
book.set_value(book.sheet("Places"), 3, 2, "Hard to find")
book.done()
assert w.save()
w.close()
app.processEvents()

wb2 = openpyxl.load_workbook(path)
assert wb2["Places"]["B3"].value == "Hard to find", wb2["Places"]["B3"].value
assert wb2["Fauna"]["A2"].value == "Sand hare"
assert "A5:B5" in [str(r) for r in wb2["Places"].merged_cells.ranges]
w2 = nousatlas.MainWindow(); w2.show()
assert w2.open_path(path)
w2.book.dirty = False

# Help > Open the sample world: copies it into "Documents" (a temp folder here) and opens the copy
docs = os.path.join(here, "Documents")
nousatlas.QStandardPaths.writableLocation = staticmethod(lambda *_: docs)
assert w2.open_sample(first_run=True)
copy = os.path.join(docs, "Gullwing Isle (sample).xlsx")
assert os.path.exists(copy) and w2.book.path == copy
book = w2.book
assert list(book.meta["maps"]) == ["Island Map"]
assert [g["name"] for g in book.meta["groups"]] == ["Welcome", "The island", "Notes"]
assert [l.name for l in book.ranges.layers] == ["Puffins", "Glow moths", "Red fox", "Grey seal", "Silver Brook"]
assert book.ranges.scheme.preset == "daynight"
w2.open_sheet("Island Map")
app.processEvents()
card = w2.squares.card("Island Map", (5, 8))      # E8, the lighthouse
assert card and "Gullwing Lighthouse" in str(card), card
w2.close()
print("OK smoke")
