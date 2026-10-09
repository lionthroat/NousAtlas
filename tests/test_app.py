r"""The whole window on a copy of the sample workbook: sidebar, editing, undo,
rows, find, range layers, save and reopen.

Run from the project folder:  .venv\Scripts\python.exe tests\test_app.py
Opens a real (offscreen) window; prints OK at the end."""
import os, sys
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("QT_QPA_FONTDIR", "C:/Windows/Fonts")
os.environ["NOUSATLAS_SETTINGS"] = "NousAtlasTest"
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import sample_copy
from PySide6.QtCore import Qt
import faulthandler
faulthandler.dump_traceback_later(120, exit=True)      # never hang on a stray dialog
from PySide6.QtWidgets import QApplication, QInputDialog, QMessageBox
import openpyxl
import nousatlas, atlas_ranges as R

QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.Yes)
QMessageBox.information = staticmethod(lambda *a, **k: None)
QMessageBox.warning = staticmethod(lambda *a, **k: None)
QInputDialog.getItem = staticmethod(lambda parent, title, label, items, *a, **k: (items[0], True))

path = sample_copy()
app = QApplication(sys.argv)
w = nousatlas.MainWindow(); w.resize(1400, 900); w.show()
assert w.open_path(path)
book = w.book

# sidebar: maps found and grouped, every sheet listed once
assert "MAP — Day One" in book.meta["maps"], book.meta["maps"]
assert book.meta["maps"]["MAP — Day One"] == {"row": 5, "col": 2, "rows": 15, "cols": 17}, book.meta["maps"]["MAP — Day One"]
assert [g["name"] for g in book.meta["groups"]] == ["Maps"]
assert len(w.sidebar_order()) == 21

# go to a sheet by typing
w.sidebar.filter.setText("fau"); w.sidebar.open_first()
assert w.ws_title == "Fauna"
w.open_sheet("Gazetteer")
assert w.history[-1][0] == "Fauna"
w.go_back(); assert w.ws_title == "Fauna"
w.go_forward(); assert w.ws_title == "Gazetteer"

# edit through the grid, then undo
g = w.grid
g.set_current(5, 5); g.start_edit("10"); g.finish_edit(True)
ws = book.sheet("Gazetteer")
assert ws["E5"].value == 10 and book.display_text(ws, 5, 7) == "5"
assert g.cur == (6, 5)
w.undo(); assert ws["E5"].value == 9
w.redo(); assert ws["E5"].value == 10

# inspector edits
g.set_current(5, 2)
assert w.inspector.text.toPlainText() == "The Springs"
w.inspector.text.setPlainText("The Springs (cold)"); w.inspector.commit_text()
assert ws["B5"].value == "The Springs (cold)"

# formatting and merge
g.select_range(5, 1, 5, 2); w.toggle_font("i")
assert ws["A5"].font.i and ws["B5"].font.i
g.select_range(2, 1, 2, 4); w.toggle_merge()
assert "A2:D2" in [str(m) for m in ws.merged_cells.ranges]
w.toggle_merge(); assert "A2:D2" not in [str(m) for m in ws.merged_cells.ranges]

# rows
g.select_range(6, 1, 6, 1); w.insert_rows(above=True)
assert ws["B7"].value == "The Horns" and ws["G7"].value.startswith("=ROUND(SQRT((E7")
w.undo(); assert ws["B6"].value == "The Horns"

# find across sheets
w.find_text("Saint's Hair")
assert len({h[0] for h in w.find_bar.hits}) >= 2, w.find_bar.hits
w.find_bar.step(1)
assert w.grid.find_current is not None

# range layers: make one from a Fauna name, paint on the Day One map
fauna = book.sheet("Fauna")
w.open_sheet("Fauna"); g.set_current(3, 1)
name = fauna["A3"].value
w.make_layer_for(name)
assert w.ws_title == "MAP — Day One" and w.range_state.painting and w.range_state.solo == name

class Ev:
    def __init__(self, mods=Qt.NoModifier): self.m = mods
    def button(self): return Qt.LeftButton
    def modifiers(self): return self.m

info = book.meta["maps"]["MAP — Day One"]
brush = w.grid.tool
w.ranges_panel.set_blocks(R.DAY)
r, c = R.square_cell(info, (9, 6))                 # I6
assert brush.press(r, c, Ev())
brush.move(r, c + 1, Ev()); brush.move(r, c + 2, Ev())
brush.release(Ev())
layer = book.ranges.get(name)
assert layer.level("MAP — Day One", 8, (9, 6)) == 3 and layer.level("MAP — Day One", 8, (11, 6)) == 3
assert layer.level("MAP — Day One", 20, (9, 6)) == 0
w.ranges_panel.set_blocks({20}); w.range_state.level = 2
assert brush.press(*R.square_cell(info, (1, 1)), Ev()); brush.release(Ev())
assert layer.level("MAP — Day One", 20, (1, 1)) == 2
who = book.ranges.who("MAP — Day One", (9, 6))
assert who and who[0][0].name == name
w.undo(); assert layer.level("MAP — Day One", 20, (1, 1)) == 0 or book.ranges.get(name).level("MAP — Day One", 20, (1, 1)) == 0
w.redo()

# a chimney-bat style layer: common at home from 16:00, uncommon around it, everywhere from 20:00
w.range_state.painting = False
bats = book.ranges.add("Chimney Bats")
book.ranges.paint(bats, "MAP — Day One", {16}, [(5, 3)], 3)
book.ranges.paint(bats, "MAP — Day One", {16}, [(4, 2), (5, 2), (6, 2), (4, 3), (6, 3), (4, 4), (5, 4), (6, 4)], 2)
book.ranges.paint(bats, "MAP — Day One", {20, 0, 4}, [(x, y) for x in range(1, 18) for y in range(1, 16)], 3)
book.dirty = True

# paint into a screenshot to make sure drawing doesn't fail
w.range_state.split = True; w.range_state.solo = None
w.right.setCurrentIndex(1); app.processEvents()
shot_dir = os.path.join(os.path.dirname(path), "shots"); os.makedirs(shot_dir, exist_ok=True)
w.grab().save(os.path.join(shot_dir, "ranges.png"))

# save, reopen: everything back
assert w.save()
wb = openpyxl.load_workbook(path)
assert "Ranges" in wb.sheetnames and wb["_NousAtlas"].sheet_state == "veryHidden"
rows = list(wb["Ranges"].iter_rows(values_only=True))
assert rows[0] == tuple(R.HEADERS), rows[0]
bat_rows = [r[:6] for r in rows if r[0] == "Chimney Bats"]
assert ("Chimney Bats", "MAP — Day One", "Night", "common", "A1:Q15", bats.color) in bat_rows, bat_rows
assert any(r[2] == "16:00" and r[3] == "common" and r[4] == "E3" for r in bat_rows), bat_rows
assert any(r[2] == "16:00" and r[3] == "uncommon" and R.parse_squares(r[4]) == {(4, 2), (5, 2), (6, 2), (4, 3), (6, 3), (4, 4), (5, 4), (6, 4)} for r in bat_rows), bat_rows
assert len(bat_rows) == 3, bat_rows
assert wb["Gazetteer"]["E5"].value == 10

w.book.dirty = False
w.close_book()
assert w.open_path(path)
b2 = w.book
assert b2.ranges.get("Chimney Bats").level("MAP — Day One", 0, (17, 15)) == 3
assert b2.ranges.get("Chimney Bats").level("MAP — Day One", 16, (5, 2)) == 2
assert b2.ranges.get(name).level("MAP — Day One", 12, (10, 6)) == 3
assert [g["name"] for g in b2.meta["groups"]] == ["Maps"]
assert "Ranges" not in [ws.title for ws in w.sidebar_order()]      # Atlas writes it; hidden from the sidebar
print("OK")

# ---- resizing with the real mouse: drag a column edge, drag a row edge, double-click a row edge
from PySide6.QtTest import QTest
from PySide6.QtCore import QPoint
w.open_sheet("Fauna")
g = w.grid
g.set_current(5, 3)
app.processEvents()
vp = g.viewport()
ws = w.book.sheet("Fauna")
# column C's right edge in the header
x = g.x_of(3) + g.colx[3] - g.colx[2] - 1
y = g.header_h // 2
before = ws.column_dimensions["C"].width
QTest.mousePress(vp, Qt.LeftButton, Qt.NoModifier, QPoint(x, y))
QTest.mouseMove(vp, QPoint(x + 60, y))
QTest.mouseRelease(vp, Qt.LeftButton, Qt.NoModifier, QPoint(x + 60, y))
app.processEvents()
after = ws.column_dimensions["C"].width
assert after > before + 5, (before, after)
assert g.colx[3] - g.colx[2] > 300          # and the grid shows it
# row 6's bottom edge in the row header: drag it taller than compact allows
yb = g.y_of(6) + g.rowy[6] - g.rowy[5] - 1
QTest.mousePress(vp, Qt.LeftButton, Qt.NoModifier, QPoint(g.header_w // 2, yb))
QTest.mouseMove(vp, QPoint(g.header_w // 2, yb + 120))
QTest.mouseRelease(vp, Qt.LeftButton, Qt.NoModifier, QPoint(g.header_w // 2, yb + 120))
app.processEvents()
assert g.rowy[6] - g.rowy[5] > 150, g.rowy[6] - g.rowy[5]
# double-click row 5's bottom edge: fits all its text (F5 is very long)
yb = g.y_of(5) + g.rowy[5] - g.rowy[4] - 1
QTest.mouseDClick(vp, Qt.LeftButton, Qt.NoModifier, QPoint(g.header_w // 2, yb))
app.processEvents()
assert g.rowy[5] - g.rowy[4] > 400, g.rowy[5] - g.rowy[4]
assert 5 in w.book.meta["pinned_rows"]["Fauna"] and 6 in w.book.meta["pinned_rows"]["Fauna"]
w.undo(); app.processEvents()
assert g.rowy[5] - g.rowy[4] < 100
print("OK resize")

# ---- per-sheet margin, cell padding, cell spacing
import atlas_dialogs as D
w.open_sheet("Fauna"); app.processEvents()
g = w.grid
plain_x = g.x_of(1)
dlg = D.SheetLayoutDialog(w)
dlg.set_values((24, 10, 6, 3))           # live preview
assert g.sheet_layout()["margin"] == 24
dlg.accept()
g._layout()
assert g.x_of(1) == g.header_w + 24 and g.y_of(1) == g.header_h + 24
rect = g.cell_rect(4, 1)
assert rect.width() == g.colx[1] - g.colx[0] - 3          # the gap shows the sheet
assert g.hit(QPoint(g.header_w + 5, g.y_of(6) + 2)) == (6, 1)   # a click in the margin lands on column A
assert ws.column_dimensions["A"].width == book.sheet("Fauna").column_dimensions["A"].width
w.grab().save(os.path.join(shot_dir, "spacing.png"))
other = w.book.sheet("Flora")
assert "Flora" not in w.book.meta["sheet_layout"]          # only this sheet
w.undo(); g._layout()
assert g.x_of(1) == plain_x
w.redo(); g._layout()
assert g.x_of(1) == g.header_w + 24
w.rename_sheet = None
print("OK spacing")

# ---- font dropdown: fonts this workbook uses come first
fb = w.font_box
fb.fill()
labels = [fb.itemText(i) for i in range(fb.count())]
assert labels[0] == "In this workbook", labels[:5]
used_end = labels.index("All fonts")
assert "Arial" in labels[1:used_end] and labels.index("Arial") < used_end
w.open_sheet("Flora"); w.grid.set_current(3, 2)
idx = fb.findData("Georgia") if fb.findData("Georgia") >= 0 else fb.findData(fb.families[0])
name = fb.itemData(idx)
fb.activated.emit(idx)
assert w.book.sheet("Flora").cell(3, 2).font.name == name
fb.fill()
assert name in [fb.itemText(i) for i in range(1, [fb.itemText(i) for i in range(fb.count())].index("All fonts"))]
fb.setEditText("times new roman"); fb.lineEdit().returnPressed.emit()
assert w.book.sheet("Flora").cell(3, 2).font.name == "Times New Roman"
print("OK fonts")

# ---- frozen panes: unfreeze, refreeze, undo
w.open_sheet("Gazetteer"); g = w.grid
gz = w.book.sheet("Gazetteer")
assert gz.freeze_panes == "C5"
w.set_freeze(None); g._layout()
assert gz.freeze_panes is None and g.fcol == 1 and g.frow == 1
w.undo(); g._layout(); assert gz.freeze_panes == "C5" and g.fcol == 3
g.set_current(5, 2); w.set_freeze("here"); assert gz.freeze_panes == "B5"
w.set_freeze("row"); assert gz.freeze_panes == "A2"
assert w.font_box.minimumWidth() >= 200
print("OK freeze")

# ---- editing a cell whose text spills: the editor shows all of it
from PySide6.QtWidgets import QMenu, QToolButton
w.open_sheet("Gazetteer"); g = w.grid
g.set_current(1, 1); g.start_edit(); app.processEvents()
assert g.editor.toPlainText().startswith("GAZETTEER")
assert g.editor.width() > g.cell_rect(1, 1).width() * 3, (g.editor.width(), g.cell_rect(1, 1).width())
g.finish_edit(False)

# ---- links as buttons: link the Cell column to the Day One map
gz = w.book.sheet("Gazetteer")
g.select_range(5, 4, 9, 4)
w.link_to_square("MAP — Day One")
assert gz["D5"].hyperlink.location == "'MAP — Day One'!J5", gz["D5"].hyperlink.location   # I1
assert gz["D9"].hyperlink.location == "'MAP — Day One'!I14"                                  # H10
g.set_current(3, 1); app.processEvents()
g.ensure_visible(5, 4); app.processEvents()
pill = g.link_rect(5, 4)
QTest.mouseClick(g.viewport(), Qt.LeftButton, Qt.NoModifier, pill.center())
app.processEvents()
assert w.ws_title == "MAP — Day One" and g.cur == (5, 10), (w.ws_title, g.cur)
w.go_back(); assert w.ws_title == "Gazetteer"
# clicking beside the button selects the cell instead
rect = g.cell_rect(5, 4); pill = g.link_rect(5, 4)
beside = QPoint(rect.left() + 2, rect.top() + 2)
assert not pill.contains(beside)
QTest.mouseClick(g.viewport(), Qt.LeftButton, Qt.NoModifier, beside)
assert w.ws_title == "Gazetteer" and g.cur == (5, 4)
# the right-click menu offers follow / remove; Fauna names offer a range layer
m = QMenu(); w.grid_menu(m)
texts = [a.text() for a in m.actions()]
assert "Follow link" in texts and "Remove link (keep the text)" in texts, texts
w.remove_links(); assert gz["D5"].hyperlink is None and gz["D5"].value == "I1"
w.undo(); assert gz["D5"].hyperlink is not None
w.open_sheet("Fauna"); g.set_current(6, 1)
m = QMenu(); w.grid_menu(m)
assert any(a.text().startswith("Make a layer for") for a in m.actions()), [a.text() for a in m.actions()]
assert w.save()
wb = openpyxl.load_workbook(path)
assert wb["Gazetteer"]["D5"].hyperlink.location == "'MAP — Day One'!J5"
print("OK links")

# ---- square cards: places from links, life from range layers, notes; hover; click through
import atlas_squares as SQ
day = "MAP — Day One"
gz = w.book.sheet("Gazetteer")
w.open_sheet("Gazetteer"); w.grid.select_range(5, 4, 20, 4); w.link_to_square(day)
row_m6 = next(r for r in range(5, 25) if gz.cell(r, 1).value == "m6")
assert gz.cell(row_m6, 4).value.startswith("G6")
fauna = w.book.sheet("Fauna")
jay_row = next(r for r in range(1, 40) if fauna.cell(r, 1).value == "Tally jay")
w.book.begin([]); jays = w.book.ranges.add("Tally jay"); w.book.ranges.paint(jays, day, set(R.DAY), [(7, 6)], 3); w.book.done("ranges")
info = w.book.meta["maps"][day]
cd = w.squares.card(day, (7, 6))
assert any(lbl.startswith("m6") and "Scout's Transect" in lbl for _, _, _, lbl in cd["places"]), cd["places"]
jay = [e for e in cd["life"] if e[0].name == "Tally jay"]
assert jay and jay[0][3] == ("Fauna", jay_row, 1) and jay[0][2] == R.DAY, cd["life"]
html_text = SQ.card_html(cd, w.theme, day)
assert "Fauna: " in html_text and "Scout&#x27;s Transect" in html_text, html_text
# note from the Cell panel
w.open_sheet(day); r6, c6 = R.square_cell(info, (7, 6)); w.grid.set_current(r6, c6)
assert w.inspector.square_box.isVisible() or not w.inspector.isVisible()
w.inspector.square_note.setPlainText("The posts are metal; the jays love them."); w.inspector.commit_square_note()
assert w.squares.card(day, (7, 6))["note"].startswith("The posts")
# hover shows the card; clicking the jay goes to its Fauna row
w.grid.ring_selection
w.on_hover(r6, c6, w.grid.viewport().mapToGlobal(w.grid.cell_rect(r6, c6).center()))
w.show_card(); app.processEvents()
assert w.card.isVisible() and w.card.key == (day, (7, 6))
w.card.clicked(SQ.href("Fauna", jay_row, 1))
assert w.ws_title == "Fauna" and w.grid.cur == (jay_row, 1) and not w.card.isVisible()
w.go_back(); assert w.ws_title == day
# renaming the map keeps the links pointing at it
w.book.begin(); w.book.rename_sheet(w.book.sheet(day), "Redsaint Map"); w.book.done("sheets")
w.ws_title = "Redsaint Map" if w.ws_title == day else w.ws_title
assert gz["D5"].hyperlink.location == "'Redsaint Map'!J5", gz["D5"].hyperlink.location
assert w.squares.card("Redsaint Map", (7, 6))["note"].startswith("The posts")
assert w.save()
wb = openpyxl.load_workbook(path)
print("OK squares")

# ---- painting has a way out; layers show only on the Layers tab; notes from the right-click menu
mp = "Redsaint Map"
w.open_sheet("Fauna"); w.grid.set_current(6, 1)
name6 = w.book.sheet("Fauna").cell(6, 1).value
w.make_layer_for(name6)
assert w.range_state.painting and w.paint_bar.isVisibleTo(w) and w.layers_showing()
QTest.keyClick(w.grid, Qt.Key_Escape)
assert not w.range_state.painting and not w.paint_bar.isVisibleTo(w)
w.ranges_panel.paint_btn.setChecked(True); assert w.range_state.painting and w.paint_bar.isVisibleTo(w)
w.right.setCurrentIndex(0)
assert not w.range_state.painting and not w.layers_showing()
w.ranges_panel.paint_btn.setChecked(False)
w._note_dialog = lambda *a: "The Tally Post: a pillar with a dish."
info = w.book.meta["maps"][mp]
w.open_sheet(mp); w.grid.set_current(*R.square_cell(info, (10, 4)))
m = QMenu(); w.grid_menu(m)
act = next(a for a in m.actions() if a.text().startswith("Note for square J4"))
act.trigger()
assert w.squares.card(mp, (10, 4))["note"].startswith("The Tally Post")
w._note_dialog = lambda *a: "cell note"
w.open_sheet("Gazetteer"); w.grid.set_current(3, 2); w.a_note.trigger()
assert w.book.sheet("Gazetteer").cell(3, 2).comment.text == "cell note"
print("OK paint exits")

# ---- Ctrl+L: layer on the selected squares; toolbar button
QInputDialog.getItem = staticmethod(lambda parent, title, label, items, cur, editable=True, *a, **k: ("Dust devils", True))
w.right.setCurrentIndex(0)
w.open_sheet(mp)
ra, ca = R.square_cell(info, (2, 2)); rb, cb = R.square_cell(info, (3, 3))
w.grid.select_range(ra, ca, rb, cb)
QTest.keyClick(w, Qt.Key_L, Qt.ControlModifier)
dd = w.book.ranges.get("Dust devils")
assert dd is not None and dd.level(mp, 12, (3, 3)) == 3 and dd.level(mp, 12, (2, 2)) == 3
assert w.layers_showing() and w.layers_btn.isChecked()
w.layers_btn.click(); assert not w.layers_showing()
w.undo(); assert w.book.ranges.get("Dust devils") is None
print("OK ctrl+l")

# ---- Here table: edit what's on a square, by time, without repainting
w.open_sheet(mp); w.right.setCurrentIndex(1)
jays = w.book.ranges.get("Tally jay")
rj, cj = R.square_cell(info, (7, 6)); w.grid.set_current(rj, cj); app.processEvents()
p = w.ranges_panel
assert "Tally jay" in p.here_layers, p.here_layers
row = p.here_layers.index("Tally jay")
col20 = 1 + R.BLOCKS.index(20)
assert jays.level(mp, 20, (7, 6)) == 0
p.here_table.cellClicked.emit(row, col20); assert jays.level(mp, 20, (7, 6)) == 3      # none -> common
row = p.here_layers.index("Tally jay")
p.here_table.cellClicked.emit(row, col20); assert jays.level(mp, 20, (7, 6)) == 2      # -> uncommon
row = p.here_layers.index("Tally jay")
p.here_table.cellClicked.emit(row, 7)                                                    # All: common -> uncommon everywhere
assert all(jays.level(mp, b, (7, 6)) == 2 for b in R.BLOCKS)
w.undo(); assert w.book.ranges.get("Tally jay").level(mp, 8, (7, 6)) == 3 and w.book.ranges.get("Tally jay").level(mp, 20, (7, 6)) == 2
# add another layer to this square from the dropdown (uses the Time/Paint settings)
p.set_blocks(R.NIGHT); w.range_state.level = 1
idx = p.here_add.findData("Chimney Bats")
if idx < 0:
    w.book.ranges.add("Chimney Bats"); p.update_here(); idx = p.here_add.findData("Chimney Bats")
p.here_add.activated.emit(idx)
bats = w.book.ranges.get("Chimney Bats")
assert bats.level(mp, 0, (7, 6)) in (1, 3) and "Chimney Bats" in p.here_layers
print("OK here table")

# ---- terrain from the map's own key; path layers; the dropper
info = w.book.meta["maps"][mp]
leg = w.squares.legend(mp)
assert any("Badlands" in v for v in leg.values()), leg
assert w.squares.terrain(mp, (7, 9)) and "Badlands" in w.squares.terrain(mp, (7, 9)), w.squares.terrain(mp, (7, 9))
cd = w.squares.card(mp, (7, 9))
assert "Terrain:" in SQ.card_html(cd, w.theme, mp)
# the south track as a path
w.book.begin([]); track = w.book.ranges.add("South track"); track.style = "path"
w.book.ranges.paint(track, mp, set(R.BLOCKS), [(9, y) for y in range(7, 16)], 3); w.book.done("ranges")
w.open_sheet(mp); w.right.setCurrentIndex(1); app.processEvents()
w.grab().save(os.path.join(shot_dir, "path.png"))
assert "path" in SQ.card_html(w.squares.card(mp, (9, 10)), w.theme, mp)
# dropper: pick the badlands red, put it on another square
rb_, cb_ = R.square_cell(info, (7, 9))
red = model_hex = None
import atlas_model as M
red = M.resolve_color(w.book.sheet(mp).cell(rb_, cb_).fill.fgColor, w.book.theme)
QTest.mouseClick(w.grid.viewport(), Qt.LeftButton, Qt.AltModifier, w.grid.cell_rect(rb_, cb_).center())
assert w.last_fill_color == red
tr_, tc_ = R.square_cell(info, (9, 10))
w.grid.set_current(tr_, tc_)
QTest.keyClick(w, Qt.Key_F, Qt.ControlModifier | Qt.ShiftModifier)
assert M.resolve_color(w.book.sheet(mp).cell(tr_, tc_).fill.fgColor, w.book.theme) == red
assert "Badlands" in w.squares.terrain(mp, (9, 10))
assert w.save()
wb = openpyxl.load_workbook(path)
assert any(r[0] == "South track" and r[6] == "path" for r in wb["Ranges"].iter_rows(values_only=True))
w.close_book(); w.open_path(path)
assert w.book.ranges.get("South track").style == "path"
print("OK terrain/path/dropper")

# ---- Delete on whole rows / columns deletes them; on cells it clears
w.open_sheet("Flora"); g = w.grid; fl = w.book.sheet("Flora")
a3, a4 = fl["A3"].value, fl["A4"].value
y = g.y_of(3) + 3
QTest.mouseClick(g.viewport(), Qt.LeftButton, Qt.NoModifier, QPoint(g.header_w // 2, y))   # row header
assert g.sel_kind == "rows"
QTest.keyClick(g, Qt.Key_Delete)
assert fl["A3"].value == a4, (fl["A3"].value, a4)
w.undo(); assert fl["A3"].value == a3
g.set_current(3, 2); QTest.keyClick(g, Qt.Key_Space, Qt.ControlModifier)
assert g.sel_kind == "cols"
b1 = fl["C3"].value
QTest.keyClick(g, Qt.Key_Delete); assert fl["B3"].value == b1
w.undo()
g.select_range(3, 1, 3, fl.max_column)          # dragged across cells: only clears
QTest.keyClick(g, Qt.Key_Delete)
assert fl["A3"].value is None and fl["A4"].value == a4
w.undo(); assert fl["A3"].value == a3
print("OK delete")

# ---- dropper picks "no fill"; size steppers; top border; border colour recolours
import atlas_dialogs as Dg
w.open_sheet("Flora"); g = w.grid; fl = w.book.sheet("Flora")
w.pick_fill_from(30, 30)                       # an empty cell: no fill
assert w.last_fill_color == ""
fl["B4"].fill = M.PatternFill("solid", fgColor="FFC0604A")
g.set_current(4, 2); w.fill_color_btn.click()
assert not fl["B4"].fill.fill_type, fl["B4"].fill.fill_type
size = fl["B4"].font.sz or 11
w.step_font_size(1); assert fl["B4"].font.sz == round(size) + 1
w.step_font_size(-1); w.step_font_size(-1); assert fl["B4"].font.sz == round(size) - 1
g.select_range(4, 1, 5, 3); w.apply_borders("top")
assert fl["A4"].border.top.style == "thin" and fl["A5"].border.top is None or not fl["A5"].border.top.style
Dg.ColorPicker.get = staticmethod(lambda *a, **k: "88C0D0")
w.choose_border_color()
assert fl["B4"].border.top.color.rgb.endswith("88C0D0"), fl["B4"].border.top.color
assert w.toolbar.findChildren(QToolButton) and w.spacing_btn.isVisibleTo(w)
assert w.a_clear_fmt not in w.toolbar.actions()
print("OK toolbar")

# ---- Layers panel can't get lost; Draw-as from the button menu; only the selected area draws
w.open_sheet(mp)
w.right.setCurrentIndex(0)
w.splitter.setSizes([200, 1200, 0]); app.processEvents()
w.layers_btn.click(); app.processEvents()
assert w.splitter.sizes()[2] >= 260 and w.layers_showing(), w.splitter.sizes()
w.range_state.layer = "South track"
w.sync_style_menu(); assert w.style_menu.title() == "Draw South track as"
w.set_layer_style("marker"); assert w.book.ranges.get("South track").style == "marker"
w.undo(); assert w.book.ranges.get("South track").style == "path"
class _P:                       # records what the overlay would draw
    def __init__(self): self.fills = 0
    def save(self): pass
    def restore(self): pass
    def setRenderHint(self, *a): pass
    def setPen(self, *a): pass
    def setBrush(self, *a): pass
    def fillRect(self, *a): self.fills += 1
    def drawPolygon(self, *a): self.fills += 1
    def drawLine(self, *a): pass
    def drawPoint(self, *a): pass
    def drawEllipse(self, *a): pass
st = w.range_state; st.solo = None; st.split = False; st.blocks = set(R.BLOCKS)
fs = w.book.ranges.get("Tally jay")
r7, c7 = R.square_cell(info, (7, 6))
st.layer = "South track"; st.all_areas = False
pp = _P(); w.grid.overlay.paint(pp, w.book.sheet(mp), r7, c7, w.grid.cell_rect(r7, c7)); assert pp.fills == 0
st.layer = "Tally jay"
pp = _P(); w.grid.overlay.paint(pp, w.book.sheet(mp), r7, c7, w.grid.cell_rect(r7, c7)); assert pp.fills == 1
st.all_areas = True
pp = _P(); w.grid.overlay.paint(pp, w.book.sheet(mp), r7, c7, w.grid.cell_rect(r7, c7)); assert pp.fills >= 1
st.all_areas = False
QTest.keyPress(w.grid, Qt.Key_Alt); assert w.grid.viewport().cursor().shape() == Qt.CrossCursor
QTest.keyRelease(w.grid, Qt.Key_Alt); assert w.grid.viewport().cursor().shape() != Qt.CrossCursor
print("OK layers panel")
