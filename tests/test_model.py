r"""Model: open, insert/delete rows and columns, undo/redo, save round trip.

Run from the project folder:  .venv\Scripts\python.exe tests\test_model.py"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import sample_copy
import openpyxl
from atlas_model import Book

path = sample_copy()
book = Book.open(path)
assert book.warnings == [], book.warnings
gz = book.wb["Gazetteer"]
assert book.display_text(gz, 5, 7) == "5", book.display_text(gz, 5, 7)
merges_before = len(book.wb["Finds & Caches"].merged_cells.ranges)
h7 = gz.row_dimensions[7].height

book.begin(); book.insert(gz, "row", 6, 2); book.done()
assert gz["G8"].value == "=ROUND(SQRT((E8-9)^2+(F8-6)^2),0)", gz["G8"].value
assert gz["B8"].value == "The Horns", gz["B8"].value
assert gz.row_dimensions[9].height == h7
assert book.display_text(gz, 8, 7) == "4"
assert gz.freeze_panes == "C5"

book.begin(); book.insert(gz, "col", 1, 1); book.done()
assert gz["H5"].value == "=ROUND(SQRT((F5-9)^2+(G5-6)^2),0)", gz["H5"].value
assert gz.freeze_panes == "D5", gz.freeze_panes

fc = book.wb["Finds & Caches"]
book.begin(); book.delete(fc, "row", 1, 3); book.done()
assert len(fc.merged_cells.ranges) == merges_before

assert book.undo() and book.undo() and book.undo()
assert gz["G5"].value.startswith("=ROUND(SQRT((E5") and gz["B6"].value == "The Horns"
assert gz.freeze_panes == "C5"
assert book.redo()
assert gz["G8"].value.startswith("=ROUND(SQRT((E8")

book.begin([gz]); book.set_value(gz, 5, 5, "10"); book.done()
assert book.display_text(gz, 5, 7) == "5"   # (10-9, 1-6) -> sqrt(26) = 5.1 -> 5
book.rename_sheet(gz, "Places")
book.save()
wb2 = openpyxl.load_workbook(path)
assert "Places" in wb2.sheetnames and "_NousAtlas" in wb2.sheetnames
assert wb2["_NousAtlas"].sheet_state == "veryHidden"
assert wb2["Places"]["E5"].value == 10
book2 = Book.open(path)
assert book2.wb["Places"].freeze_panes == "C5"
print("OK")
